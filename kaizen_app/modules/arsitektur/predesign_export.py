"""Excel and PDF reports for the Pre-Design module."""
from __future__ import annotations

from dataclasses import fields
from datetime import datetime
from html import escape
from io import BytesIO
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    KeepTogether,
    Image as PdfImage,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

_NAVY = "1F3864"
_PALE = "EEF2F8"
_YELLOW = "FFF200"
_LINE = "8C96A3"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_PDF_NAVY = colors.HexColor(f"#{_NAVY}")
_PDF_PALE = colors.HexColor(f"#{_PALE}")
_PDF_LINE = colors.HexColor(f"#{_LINE}")


def _display(value: Any) -> str:
    if value is None or value == "":
        return "—"
    if isinstance(value, bool):
        return "Ya" if value else "Tidak"
    if isinstance(value, dict):
        return "; ".join(f"{key}: {_display(item)}" for key, item in value.items())
    if isinstance(value, (list, tuple)):
        return ", ".join(_display(item) for item in value) or "—"
    return str(value)


def _input_rows(model: Any) -> list[tuple[str, Any]]:
    return [
        (field.name.replace("_", " ").capitalize(), getattr(model, field.name))
        for field in fields(model)
    ]


def _floor_assignments(project) -> dict[str, list[int]]:
    assignments: dict[str, list[int]] = {}
    for floor, items in sorted(project.floor_programs.items()):
        for item in items:
            assignments.setdefault(item.room_id, []).append(floor)
    return assignments


def _style_sheet(ws, widths: list[int]) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width


def _table_header(ws, row: int, start: int, end: int) -> None:
    for cells in ws.iter_rows(min_row=row, max_row=row, min_col=start, max_col=end):
        for cell in cells:
            cell.fill = PatternFill("solid", fgColor=_NAVY)
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color=_LINE))


def _append_rows(ws, rows: list[list[Any]]) -> None:
    for row in rows:
        ws.append([_display(value) for value in row])
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color=_LINE))


def build_excel(project, findings, completeness, envelope, program_result, *, layout_images=None) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Ringkasan"
    _style_sheet(summary, [34, 70, 20, 20])
    summary.merge_cells("A1:D1")
    summary["A1"] = f"LAPORAN PRE-DESIGN {project.building_type.upper()}"
    summary["A1"].fill = PatternFill("solid", fgColor=_NAVY)
    summary["A1"].font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    summary["A1"].alignment = Alignment(vertical="center")
    summary.row_dimensions[1].height = 28
    _append_rows(summary, [
        ["Nama proyek", project.name],
        ["Lokasi", project.location],
        ["Submodel bangunan", project.building_type],
        ["Tanggal sesi", project.session_date],
        ["Tanggal laporan", datetime.now().strftime("%Y-%m-%d %H:%M")],
        [],
        ["Ringkasan validasi", "Jumlah"],
        ["BLOCKER", sum(item.severity.value == "BLOCKER" for item in findings)],
        ["WARNING", sum(item.severity.value == "WARNING" for item in findings)],
        ["INFO", sum(item.severity.value == "INFO" for item in findings)],
    ])
    summary["A8"].fill = PatternFill("solid", fgColor=_YELLOW)
    summary["B8"].fill = PatternFill("solid", fgColor=_YELLOW)
    summary["A8"].font = Font(bold=True)
    summary["B8"].font = Font(bold=True)
    summary.append([])
    summary.append(["Kelengkapan form", "Terisi", "Persentase", "Keputusan"])
    _table_header(summary, summary.max_row, 1, 4)
    for section_id, section in completeness.items():
        summary.append([
            f"{section_id} · {section['title']}",
            f"{section['filled']}/{section['required']}",
            f"{section['pct']:.0f}%",
            "Ya" if section["decided"] else "Belum",
        ])
    if envelope:
        summary.append([])
        summary.append(["Analisis intensitas bangunan", "Hasil"])
        _table_header(summary, summary.max_row, 1, 2)
        _append_rows(summary, [
            ["Luas lahan (m²)", envelope.land_area_m2],
            ["Batas tapak KDB (m²)", envelope.footprint_cap_m2],
            ["Batas GFA KLB (m²)", envelope.gfa_cap_m2],
            ["KDH minimum (m²)", envelope.min_green_m2],
            ["Lantai efektif", envelope.floors_effective],
        ])
    if program_result:
        summary.append([])
        summary.append(["Estimasi program ruang", "Hasil"])
        _table_header(summary, summary.max_row, 1, 2)
        _append_rows(summary, [
            ["Estimasi GFA (m²)", program_result.gfa_estimate_m2],
            ["Estimasi footprint (m²)", program_result.footprint_estimate_m2],
            ["Luas tertutup bersih (m²)", program_result.net_enclosed_m2],
            ["Jumlah lantai asumsi", program_result.floors_assumed],
        ])

    parameters = workbook.create_sheet("Parameter")
    _style_sheet(parameters, [30, 38, 60])
    parameters.append(["Seksi", "Parameter", "Nilai"])
    _table_header(parameters, 1, 1, 3)
    for section, model in (
        ("Pembeli", project.persona),
        ("Produk", project.product),
        ("Regulasi", project.regulation),
        ("Tapak", project.site),
        ("Fasad", project.facade),
    ):
        for label, value in _input_rows(model):
            parameters.append([section, label, _display(value)])
    for row in parameters.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color=_LINE))

    rooms = workbook.create_sheet("Program Ruang")
    _style_sheet(rooms, [12, 32, 15, 12, 16, 16, 24, 14])
    rooms.append([
        "Lantai", "Ruang", "Prioritas", "Jumlah", "Luas/unit (m²)", "Total (m²)",
        "Status", "Di bawah minimum",
    ])
    _table_header(rooms, 1, 1, 8)
    if program_result:
        floor_assignments = _floor_assignments(project)
        assignment_index: dict[str, int] = {}
        for line in program_result.lines:
            index = assignment_index.get(line.room_id, 0)
            floors = floor_assignments.get(line.room_id, [])
            floor = floors[index] if index < len(floors) else ""
            assignment_index[line.room_id] = index + 1
            rooms.append([
                floor, line.name, line.priority, line.qty, line.unit_area_m2,
                line.total_m2, line.status, "Ya" if line.below_min else "Tidak",
            ])
    if rooms.max_row == 1:
        rooms.append(["Belum ada program ruang"])
    if program_result:
        rooms.append([])
        rooms.append(["Estimasi GFA (m²)", program_result.gfa_estimate_m2])
        rooms.append(["Estimasi footprint (m²)", program_result.footprint_estimate_m2])
        rooms.append(["Luas bersih tertutup (m²)", program_result.net_enclosed_m2])
    for row in rooms.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color=_LINE))

    validation = workbook.create_sheet("Validasi")
    _style_sheet(validation, [14, 14, 16, 64, 42, 64])
    validation.append(["Severity", "Rule", "Seksi", "Temuan", "Rujukan", "Saran"])
    _table_header(validation, 1, 1, 6)
    for finding in findings:
        validation.append([
            finding.severity.value, finding.rule_id, finding.section,
            finding.message, finding.reference, finding.suggestion,
        ])
    for row in validation.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color=_LINE))
    if validation.max_row == 1:
        validation.append(["Tidak ada temuan"])

    if layout_images:
        layout_sheet = workbook.create_sheet("Simulasi Ruang")
        layout_sheet.sheet_view.showGridLines = False
        layout_sheet.column_dimensions["A"].width = 125
        layout_sheet["A1"] = "SIMULASI ZONING KONSEP"
        layout_sheet["A1"].font = Font(name="Arial", size=14, bold=True, color=_NAVY)
        layout_sheet["A2"] = (
            "Skema otomatis, tidak berskala, dan bukan pengganti gambar kerja arsitek."
        )
        row = 4
        for floor_plan in layout_images:
            layout_sheet.cell(row=row, column=1, value=f"Konsep zoning — lantai {floor_plan.floor}")
            layout_sheet.cell(row=row, column=1).font = Font(name="Arial", size=11, bold=True)
            image = ExcelImage(BytesIO(floor_plan.png))
            scale = 900 / image.width
            image.height = round(image.height * scale)
            image.width = 900
            layout_sheet.add_image(image, f"A{row + 1}")
            row += math.ceil(image.height / 20) + 3

    decisions = workbook.create_sheet("Keputusan")
    _style_sheet(decisions, [16, 16, 70, 26, 24])
    decisions.append(["Seksi", "Status", "Keputusan", "Diputuskan oleh", "Waktu"])
    _table_header(decisions, 1, 1, 5)
    for decision in project.decisions:
        decisions.append([
            decision.section_id,
            "Diputuskan" if decision.decided else "Belum",
            decision.decision_text,
            decision.decided_by,
            decision.decided_at,
        ])
    if decisions.max_row == 1:
        decisions.append(["Belum ada keputusan"])
    for row in decisions.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color=_LINE))

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _pdf_styles():
    base = getSampleStyleSheet()
    return (
        ParagraphStyle(
            "PreDesignTitle", parent=base["Title"], fontName="Helvetica-Bold",
            fontSize=15, leading=18, textColor=_PDF_NAVY, alignment=TA_LEFT,
            spaceAfter=4,
        ),
        ParagraphStyle(
            "PreDesignSmall", parent=base["BodyText"], fontSize=8, leading=11,
            textColor=colors.HexColor("#555555"),
        ),
        ParagraphStyle(
            "PreDesignBody", parent=base["BodyText"], fontSize=8.5, leading=11,
        ),
        ParagraphStyle(
            "PreDesignCell", parent=base["BodyText"], fontSize=7.5, leading=9,
        ),
    )


def _pdf_table(rows, widths, *, repeat_rows=1):
    table = Table(rows, colWidths=widths, repeatRows=repeat_rows, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), _PDF_NAVY),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, _PDF_LINE),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, _PDF_PALE]),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return table


def build_pdf(project, findings, completeness, envelope, program_result, *, layout_images=None) -> bytes:
    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=18 * mm,
        title=f"Laporan Pre-Design - {project.name}",
        author="Kaizen PBG",
    )
    title_style, small_style, body_style, cell_style = _pdf_styles()
    paragraph = lambda value, style=cell_style: Paragraph(escape(_display(value)), style)
    story = [
        Paragraph(f"LAPORAN PRE-DESIGN {escape(project.building_type.upper())}", title_style),
        Paragraph("Kaizen PBG · Laporan hasil analisis dan validasi", small_style),
        Spacer(1, 4 * mm),
        _pdf_table([
            [paragraph("Identitas"), paragraph("Nilai")],
            [paragraph("Nama proyek"), paragraph(project.name)],
            [paragraph("Lokasi"), paragraph(project.location)],
            [paragraph("Submodel bangunan"), paragraph(project.building_type)],
            [paragraph("Tanggal sesi"), paragraph(project.session_date)],
            [paragraph("Tanggal laporan"), paragraph(datetime.now().strftime("%Y-%m-%d %H:%M"))],
        ], [38 * mm, 136 * mm]),
        Spacer(1, 4 * mm),
        Paragraph("RINGKASAN VALIDASI", body_style),
    ]
    counts = {
        severity: sum(item.severity.value == severity for item in findings)
        for severity in ("BLOCKER", "WARNING", "INFO")
    }
    story.extend([
        _pdf_table(
            [[paragraph("BLOCKER"), paragraph("WARNING"), paragraph("INFO")],
             [paragraph(counts["BLOCKER"]), paragraph(counts["WARNING"]), paragraph(counts["INFO"])]],
            [58 * mm, 58 * mm, 58 * mm],
        ),
        Spacer(1, 4 * mm),
        Paragraph("KELENGKAPAN FORM", body_style),
    ])
    completeness_rows = [[
        paragraph("Seksi"), paragraph("Terisi"), paragraph("Kelengkapan"), paragraph("Keputusan"),
    ]]
    for section_id, section in completeness.items():
        completeness_rows.append([
            paragraph(f"{section_id} · {section['title']}"),
            paragraph(f"{section['filled']}/{section['required']}"),
            paragraph(f"{section['pct']:.0f}%"),
            paragraph("Ya" if section["decided"] else "Belum"),
        ])
    story.extend([
        _pdf_table(completeness_rows, [75 * mm, 28 * mm, 32 * mm, 39 * mm]),
        Spacer(1, 4 * mm),
        Paragraph("PARAMETER PROYEK", body_style),
    ])
    parameter_rows = [[paragraph("Seksi"), paragraph("Parameter"), paragraph("Nilai")]]
    for section, model in (
        ("Pembeli", project.persona),
        ("Produk", project.product),
        ("Regulasi", project.regulation),
        ("Tapak", project.site),
        ("Fasad", project.facade),
    ):
        parameter_rows.extend([
            [paragraph(section), paragraph(label), paragraph(value)]
            for label, value in _input_rows(model)
        ])
    story.extend([
        _pdf_table(parameter_rows, [30 * mm, 55 * mm, 89 * mm]),
        Spacer(1, 4 * mm),
    ])

    if envelope:
        story.append(Paragraph("ANALISIS INTENSITAS BANGUNAN", body_style))
        envelope_rows = [[paragraph("Parameter"), paragraph("Hasil")]]
        envelope_rows.extend([
            [paragraph(label), paragraph(value)]
            for label, value in (
                ("Luas lahan (m²)", envelope.land_area_m2),
                ("Batas tapak KDB (m²)", envelope.footprint_cap_m2),
                ("Batas GFA KLB (m²)", envelope.gfa_cap_m2),
                ("KDH minimum (m²)", envelope.min_green_m2),
                ("Lantai efektif", envelope.floors_effective),
                ("Envelope terbangun (m²)", envelope.buildable_envelope_m2),
            )
        ])
        story.extend([_pdf_table(envelope_rows, [90 * mm, 84 * mm]), Spacer(1, 4 * mm)])

    if program_result:
        story.append(Paragraph("PROGRAM RUANG", body_style))
        room_rows = [[
            paragraph("Lantai"), paragraph("Ruang"), paragraph("Prioritas"), paragraph("Qty"),
            paragraph("m²/unit"), paragraph("Total m²"), paragraph("Status"),
            paragraph("Di bawah min."),
        ]]
        floor_assignments = _floor_assignments(project)
        assignment_index: dict[str, int] = {}
        for line in program_result.lines:
            index = assignment_index.get(line.room_id, 0)
            floors = floor_assignments.get(line.room_id, [])
            floor = floors[index] if index < len(floors) else "—"
            assignment_index[line.room_id] = index + 1
            room_rows.append([
                paragraph(floor), paragraph(line.name), paragraph(line.priority), paragraph(line.qty),
                paragraph(f"{line.unit_area_m2:.1f}"), paragraph(f"{line.total_m2:.1f}"),
                paragraph(line.status), paragraph("Ya" if line.below_min else "Tidak"),
            ])
        story.extend([
            _pdf_table(room_rows, [13 * mm, 39 * mm, 21 * mm, 12 * mm, 22 * mm, 20 * mm, 25 * mm, 22 * mm]),
            Spacer(1, 2 * mm),
            Paragraph(
                escape(
                    f"Estimasi GFA: {program_result.gfa_estimate_m2:.1f} m² · "
                    f"Footprint: {program_result.footprint_estimate_m2:.1f} m² · "
                    f"Luas bersih tertutup: {program_result.net_enclosed_m2:.1f} m²"
                ),
                body_style,
            ),
            Spacer(1, 4 * mm),
        ])

    if layout_images:
        story.extend([
            Paragraph("SIMULASI TATA RUANG KONSEP", body_style),
            Paragraph(
                "Skema zoning otomatis berdasarkan zona fungsi. Tidak berskala dan bukan pengganti gambar kerja arsitek.",
                small_style,
            ),
            Spacer(1, 2 * mm),
        ])
        for floor_plan in layout_images:
            image = PdfImage(BytesIO(floor_plan.png))
            image.drawWidth = doc.width
            image.drawHeight = doc.width * image.imageHeight / image.imageWidth
            story.extend([
                Paragraph(f"Konsep zoning — lantai {floor_plan.floor}", body_style),
                image,
                Spacer(1, 4 * mm),
            ])

    story.append(Paragraph("TEMUAN VALIDASI", body_style))
    for finding in findings:
        content = [
            Paragraph(
                f"<b>{escape(finding.severity.value)} · {escape(finding.rule_id)} · "
                f"{escape(finding.section)}</b>",
                body_style,
            ),
            Paragraph(escape(finding.message), body_style),
        ]
        if finding.reference:
            content.append(Paragraph(f"Rujukan: {escape(finding.reference)}", small_style))
        if finding.suggestion:
            content.append(Paragraph(f"Saran: {escape(finding.suggestion)}", small_style))
        content.append(Spacer(1, 2 * mm))
        story.append(KeepTogether(content))
    if not findings:
        story.append(Paragraph("Tidak ada temuan validasi.", body_style))

    story.extend([Spacer(1, 2 * mm), Paragraph("KEPUTUSAN TERCATAT", body_style)])
    decision_rows = [[
        paragraph("Seksi"), paragraph("Status"), paragraph("Keputusan"),
        paragraph("Diputuskan oleh"), paragraph("Waktu"),
    ]]
    decision_rows.extend([
        [
            paragraph(item.section_id),
            paragraph("Diputuskan" if item.decided else "Belum"),
            paragraph(item.decision_text),
            paragraph(item.decided_by),
            paragraph(item.decided_at),
        ]
        for item in project.decisions
    ])
    if len(decision_rows) == 1:
        decision_rows.append([paragraph("Belum ada keputusan"), "", "", "", ""])
    story.extend([
        _pdf_table(decision_rows, [20 * mm, 25 * mm, 67 * mm, 32 * mm, 28 * mm]),
        Spacer(1, 4 * mm),
        Paragraph(
            "Nilai default/asumsi dan parameter regulasi perlu diverifikasi terhadap RDTR, "
            "ketentuan resmi, dan data lokal. Laporan ini bukan penetapan kepatuhan hukum.",
            small_style,
        ),
    ])

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(document.leftMargin, 9 * mm, "Kaizen PBG - Pre-Design")
        canvas.drawRightString(
            A4[0] - document.rightMargin,
            9 * mm,
            f"Dibuat {datetime.now():%d-%m-%Y %H:%M}  |  Halaman {document.page}",
        )
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
