"""Modul Pipa Air Hujan - ekspor Excel (rumus hidup, desain konsisten dengan tampilan web) dan PDF."""
from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from io import BytesIO

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.worksheet.pagebreak import Break
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from modules.pipa_hujan_diagram import diagram_png

NAVY, BLUE, AQUA, SOFT, LINE = "0F3D63", "2A7AB0", "35B6C9", "EAF3F9", "D5E1EB"
_THIN = Side(style="thin", color=LINE)
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_F_FILL = PatternFill("solid", fgColor=NAVY)
_SOFT = PatternFill("solid", fgColor=SOFT)
_MID = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_NW = Alignment(horizontal="left", vertical="center", wrap_text=False)
F_N = Font(name="Arial", size=10)
F_B = Font(name="Arial", size=10, bold=True)
F_IN = Font(name="Arial", size=10, color="0000FF")      # biru = input
F_H = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F_I = Font(name="Arial", size=9, italic=True, color="6B7C8C")
F_T = Font(name="Arial", size=15, bold=True, color="FFFFFF")
F_S = Font(name="Arial", size=12, bold=True, color=NAVY)
T16 = "'Tabel 16 SNI 8153'"


def _put(ws, ref, val, font=F_N, fill=None, al=_MID, fmt=None, border=True):
    c = ws[ref]
    c.value = val
    c.font, c.alignment = font, al
    if border:
        c.border = _BOX
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt


def _merge(ws, rng, val, **kw):
    ws.merge_cells(rng)
    _put(ws, rng.split(":")[0], val, **kw)
    if kw.get("border", True):
        for row in ws[rng]:
            for c in row:
                c.border = _BOX


def build_excel(h: dict, data: dict, proyek: dict) -> bytes:
    P, z = h["P"], h["zona"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Pipa Air Hujan"
    wt = wb.create_sheet("Tabel 16 SNI 8153")
    wn = wb.create_sheet("SNI & Peraturan")
    put = lambda *a, **k: _put(ws, *a, **k)
    merge = lambda *a, **k: _merge(ws, *a, **k)

    # ======================================================= sheet Tabel 16
    t16 = data["tabel16"]
    _merge(wt, "B1:J1", "SNI 8153:2015 — Tabel 16: Penentuan Ukuran Perpipaan Air Hujan Horizontal", font=F_T, fill=_F_FILL,
           al=_NW, border=False)
    wt.row_dimensions[1].height = 30
    _put(wt, "B2", "Luas bidang datar horizontal maksimum yang diperbolehkan (m²) pada berbagai curah hujan", F_I, al=_NW, border=False)
    for col, t in zip("BCD", ["Kemiringan (%)", "Ø pipa (inci)", "Debit (L/dt)"]):
        _put(wt, f"{col}4", t, F_H, _F_FILL)
    for col, g in zip("EFGHIJ", t16["intensitas"]):
        _put(wt, f"{col}4", g, F_H, _F_FILL, fmt='0.0" mm/jam"')
    r = 5
    for kem in ("1", "2", "4"):
        kd = t16["kemiringan"][kem]
        for i, u in enumerate(data["ukuran_inci"]):
            _put(wt, f"B{r}", int(kem), F_N)
            _put(wt, f"C{r}", u, F_B)
            _put(wt, f"D{r}", kd["debit"][i], F_N, fmt="0.00")
            for col, v in zip("EFGHIJ", kd["luas"][i]):
                _put(wt, f"{col}{r}", v, F_IN, fmt="#,##0")
            if int(kem) % 2 == 0:
                for col in "BCDEFGHIJ":
                    wt[f"{col}{r}"].fill = _SOFT
            r += 1
    t_end = r - 1
    _merge(wt, f"B{t_end + 2}:J{t_end + 2}", data["tabel16"]["sumber"], font=F_I, al=_LEFT, border=False)
    _merge(wt, f"B{t_end + 3}:J{t_end + 3}", "Catatan: " + data["tabel16"]["catatan_debit"], font=F_I, al=_LEFT, border=False)
    wt.row_dimensions[t_end + 2].height = 28
    wt.row_dimensions[t_end + 3].height = 40
    for col, w in {"A": 2, "B": 15, "C": 14, "D": 13, "E": 14, "F": 14, "G": 14, "H": 14, "I": 14, "J": 14}.items():
        wt.column_dimensions[col].width = w
    wt.freeze_panes = "B5"
    wt.page_setup.orientation = "landscape"
    wt.sheet_properties.pageSetUpPr.fitToPage = True
    wt.page_setup.fitToWidth, wt.page_setup.fitToHeight = 1, 1
    RNG_L = f"{T16}!$E$5:$J${t_end}"
    RNG_S = f"{T16}!$B$5:$B${t_end}"
    RNG_U = f"{T16}!$C$5:$C${t_end}"
    RNG_I = f"{T16}!$E$4:$J$4"

    # ======================================================= sheet SNI & peraturan
    _merge(wn, "B1:F1", "LANDASAN PERENCANAAN: SNI & PERATURAN TERKAIT PIPA AIR HUJAN", font=F_T, fill=_F_FILL, al=_NW, border=False)
    wn.row_dimensions[1].height = 30
    for col, t in zip("BCDEF", ["No", "Standar / Peraturan", "Judul", "Peran pada perhitungan", "Catatan"]):
        _put(wn, f"{col}3", t, F_H, _F_FILL)
    for i, s in enumerate(data["standar"]):
        rr = 4 + i
        _put(wn, f"B{rr}", i + 1)
        _put(wn, f"C{rr}", s["kode"], F_B, al=_LEFT)
        _put(wn, f"D{rr}", s["judul"], F_N, al=_LEFT)
        _put(wn, f"E{rr}", s["peran"], F_N, al=_LEFT)
        _put(wn, f"F{rr}", s["catatan"], F_N, al=_LEFT)
        wn.row_dimensions[rr].height = 44
    _put(wn, f"B{5 + len(data['standar'])}", data["catatan_standar"], F_I, al=_NW, border=False)
    for col, w in {"A": 2, "B": 5, "C": 30, "D": 44, "E": 52, "F": 40}.items():
        wn.column_dimensions[col].width = w
    wn.page_setup.orientation = "landscape"
    wn.sheet_properties.pageSetUpPr.fitToPage = True
    wn.page_setup.fitToWidth, wn.page_setup.fitToHeight = 1, 0

    # ======================================================= sheet utama
    for col, w in {"A": 2, "B": 5, "C": 36, "D": 16, "E": 16, "F": 16, "G": 18, "H": 17, "I": 18, "J": 16, "K": 4, "L": 12}.items():
        ws.column_dimensions[col].width = w
    merge("B1:J1", "PERHITUNGAN PIPA AIR HUJAN — ROOF DRAIN, PIPA HORIZONTAL DAN KOLEKTOR VERTIKAL", font=F_T, fill=_F_FILL,
          al=_NW, border=False)
    ws.row_dimensions[1].height = 34
    for r, (lab, key) in enumerate([("PEKERJAAN", "pekerjaan"), ("LOKASI", "lokasi"), ("TAHUN", "tahun"),
                                    ("ITEM PEKERJAAN", "item")], start=3):
        merge(f"B{r}:C{r}", lab, font=Font(name="Arial", size=10, bold=True, color=NAVY), fill=_SOFT, al=_NW, border=False)
        merge(f"D{r}:J{r}", f": {proyek.get(key, '')}", font=F_N, fill=_SOFT, al=_NW, border=False)
    merge("B8:J8", "Landasan — Debit: SNI 2415:2016 · Pipa air hujan: SNI 8153:2015 (Tabel 16 dan 17) · "
                   "Permen PU 11/PRT/M/2014 · Permen PU 12/PRT/M/2014 (lengkap pada sheet 'SNI & Peraturan')",
          font=F_I, al=_LEFT, border=False)
    ws.row_dimensions[8].height = 26

    r = 10

    def sec(no, judul):
        nonlocal r
        put(f"B{r}", no, F_H, _F_FILL)
        merge(f"C{r}:J{r}", judul, font=F_S, al=_NW, border=False)
        for col in "CDEFGHIJ":
            ws[f"{col}{r}"].border = Border(bottom=Side(style="medium", color=AQUA))
        ws.row_dimensions[r].height = 22
        r += 1

    # ---- skema
    sec("•", "Skema Sistem")
    img = XLImage(BytesIO(diagram_png(h, data)))
    img.width, img.height = 1030, int(1030 * 560 / 1500)
    ws.add_image(img, f"B{r}")
    r += 21

    # ---- 1. dasar perhitungan
    sec(1, "Dasar Perhitungan")
    put(f"C{r}", "Metode : Rumus Rasional (SNI 2415:2016)", F_B, al=_NW, border=False)
    r += 1
    merge(f"C{r}:G{r}", "Q = 0,00278 × C × I × A", font=Font(name="Cambria", size=13, italic=True, color=NAVY), fill=_SOFT, al=_MID)
    ws.row_dimensions[r].height = 26
    r += 1
    for a, b in [("Q", "debit air hujan (m³/dt)"), ("C", "koefisien limpasan"), ("I", "intensitas hujan (mm/jam)"),
                 ("A", "luas atap / daerah tangkapan (Ha)")]:
        put(f"C{r}", f"{a}  =  {b}", F_I, al=_NW, border=False)
        r += 1
    r += 1

    # ---- 2. data curah hujan
    sec(2, "Data Curah Hujan")
    for col, t in zip("CDE", ["Parameter", "Nilai", "Sat"]):
        put(f"{col}{r}", t, F_H, _F_FILL)
    r += 1
    put(f"C{r}", "Intensitas Hujan (I)", al=_LEFT)
    put(f"D{r}", P["I"], F_IN, fmt="0.0")
    put(f"E{r}", "mm/jam")
    RI = f"$D${r}"
    r += 1
    put(f"C{r}", "Konstanta rumus rasional", al=_LEFT)
    put(f"D{r}", data["rasional"]["konstanta"], F_IN, fmt="0.00000")
    put(f"E{r}", "–")
    RK = f"$D${r}"
    r += 1
    r_luas, r_ha, r_c = r, r + 1, r + 2
    r += 3
    r += 1
    for col, t in zip("CDEFG", ["Bidang Atap", "Luas A (m²)", "Koef. Limpasan (C)", "Jumlah Roof Drain", "Luas A (Ha)"]):
        put(f"{col}{r}", t, F_H, _F_FILL)
    ws.row_dimensions[r].height = 28
    r += 1
    z0 = r
    for x in z.itertuples():
        put(f"C{r}", x.Bidang, F_IN, al=_LEFT)
        put(f"D{r}", x.A, F_IN, fmt="#,##0.00")
        put(f"E{r}", x.C, F_IN, fmt="0.00")
        put(f"F{r}", x.n, F_IN, fmt="0")
        put(f"G{r}", f"=D{r}/10000", fmt="0.000000")
        r += 1
    z1 = r - 1
    put(f"C{r_luas}", "Luas Atap (A)", al=_LEFT)
    put(f"D{r_luas}", f"=SUM(D{z0}:D{z1})", F_B, fmt="#,##0.00")
    put(f"E{r_luas}", "m²")
    put(f"C{r_ha}", "Luas Atap (A)", al=_LEFT)
    put(f"D{r_ha}", f"=D{r_luas}/10000", F_B, fmt="0.000000")
    put(f"E{r_ha}", "Ha")
    put(f"C{r_c}", "Koef. Limpasan (C) rata-rata", al=_LEFT)
    put(f"D{r_c}", f"=SUMPRODUCT(D{z0}:D{z1},E{z0}:E{z1})/D{r_luas}", F_B, fmt="0.00")
    put(f"E{r_c}", "–")
    ATOT, CRATA = f"$D${r_luas}", f"$D${r_c}"
    r += 1

    # ---- 3. analisa teknis
    sec(3, "Analisa Perhitungan Teknis")
    for col, t in zip("CDEFG", ["Bidang Atap", "Q (m³/dt)", "Q (L/dt)", "Luas per roof drain (m²)", "Q per roof drain (L/dt)"]):
        put(f"{col}{r}", t, F_H, _F_FILL)
    ws.row_dimensions[r].height = 28
    r += 1
    a0 = r
    for i in range(len(z)):
        s = z0 + i
        put(f"C{r}", f"=C{s}", al=_LEFT)
        put(f"D{r}", f"={RK}*E{s}*{RI}*G{s}", fmt="0.00000")
        put(f"E{r}", f"=D{r}*1000", fmt="0.000")
        put(f"F{r}", f"=D{s}/F{s}", fmt="0.00")
        put(f"G{r}", f"=E{r}/F{s}", fmt="0.000")
        r += 1
    a1 = r - 1
    put(f"C{r}", "TOTAL", F_B, _SOFT, al=_LEFT)
    put(f"D{r}", f"=SUM(D{a0}:D{a1})", F_B, _SOFT, fmt="0.00000")
    put(f"E{r}", f"=SUM(E{a0}:E{a1})", F_B, _SOFT, fmt="0.000")
    put(f"F{r}", f"=MAX(F{a0}:F{a1})", F_B, _SOFT, fmt="0.00")
    put(f"G{r}", f"=MAX(G{a0}:G{a1})", F_B, _SOFT, fmt="0.000")
    QM3, QLS, ADRAIN = f"$D${r}", f"$E${r}", f"$F${r}"
    r += 1
    merge(f"C{r}:J{r}", f'="Q = "&TEXT({RK},"0.00000")&" × "&TEXT({CRATA},"0.00")&" × "&TEXT({RI},"0")&" × "&TEXT({ATOT.replace("$", "")}/10000,"0.000000")&" = "&TEXT({QM3},"0.00000")&" m³/dt = "&TEXT({QLS},"0.00")&" L/dt"',
          font=Font(name="Cambria", size=11, italic=True, color=NAVY), fill=_SOFT, al=_NW)
    r += 2

    # ---- 4. pipa horizontal
    ws.row_breaks.append(Break(id=r - 1))
    sec(4, "Pipa Horizontal / Mendatar — SNI 8153:2015 (Tabel 16)")
    put(f"C{r}", "Kolom intensitas Tabel 16 (urutan)", al=_LEFT)
    put(f"D{r}", f'=COUNTIF({RNG_I},"<"&{RI})+1', fmt="0")
    put(f"E{r}", "ke-")
    IDX = f"$D${r}"
    r += 1
    put(f"C{r}", "Intensitas kolom yang dipakai", al=_LEFT)
    put(f"D{r}", f"=IF({IDX}>6,INDEX({RNG_I},6),INDEX({RNG_I},{IDX}))", fmt="0.0")
    put(f"E{r}", "mm/jam")
    r += 2

    def cap(size, slope):
        sh = f"INDEX({RNG_L},0,{{c}})"
        return (f"IF({IDX}<=6,SUMIFS({sh.format(c=IDX)},{RNG_S},{slope},{RNG_U},{size}),"
                f"SUMIFS({sh.format(c=6)},{RNG_S},{slope},{RNG_U},{size})*{T16}!$J$4/{RI})")

    heads = ["Pipa", "Luas dilayani (m²)", "Ø (inci)", "Kemiringan (%)", "Kapasitas Tabel 16 (m²)",
             "Pemakaian kapasitas", "Status", "Ø minimum (inci)"]
    for col, t in zip("CDEFGHIJ", heads):
        put(f"{col}{r}", t, F_H, _F_FILL)
    ws.row_dimensions[r].height = 30
    r += 1
    rc, rg = r, r + 1
    r += 2
    r += 1

    # tabel kapasitas per Ø (referensi + bantu Ø minimum)
    put(f"C{r}", "Kapasitas per ukuran pipa pada kolom intensitas terpilih", F_B, al=_NW, border=False)
    r += 1
    for col, t in zip("CDEFG", ["Ø pipa (inci)", "Kapasitas — kemiringan cabang (m²)", "Kapasitas — kemiringan gabungan (m²)",
                                "Ø memenuhi (cabang)", "Ø memenuhi (gabungan)"]):
        put(f"{col}{r}", t, F_H, _F_FILL)
    ws.row_dimensions[r].height = 42
    r += 1
    u0 = r
    for u in data["ukuran_inci"]:
        put(f"C{r}", u, F_B, fmt='General"″"')
        put(f"D{r}", "=" + cap(f"$C{r}", f"$F${rc}"), fmt="#,##0")
        put(f"E{r}", "=" + cap(f"$C{r}", f"$F${rg}"), fmt="#,##0")
        put(f"F{r}", f'=IF(D{r}>=$D${rc},C{r},"")', fmt='General"″"')
        put(f"G{r}", f'=IF(E{r}>=$D${rg},C{r},"")', fmt='General"″"')
        r += 1
    u1 = r - 1
    for row, nama, luas, dia, kem, kol in [
        (rc, "Cabang (per roof drain)", f"={ADRAIN}", P["d_cabang"], P["s_cabang"], "F"),
        (rg, "Horizontal gabungan", f"={ATOT}", P["d_gabung"], P["s_gabung"], "G"),
    ]:
        put(f"C{row}", nama, F_B, al=_LEFT)
        put(f"D{row}", luas, fmt="0.00")
        put(f"E{row}", dia, F_IN, fmt='General"″"')
        put(f"F{row}", kem, F_IN, fmt="0")
        put(f"G{row}", "=" + cap(f"$E{row}", f"$F{row}"), F_B, fmt="#,##0")
        put(f"H{row}", f"=D{row}/G{row}", fmt="0%")
        put(f"I{row}", f'=IF(D{row}<=G{row},"MEMENUHI","TIDAK MEMENUHI")', F_B)
        put(f"J{row}", f'=IF(COUNT({kol}{u0}:{kol}{u1})=0,"di luar tabel",MIN({kol}{u0}:{kol}{u1}))', fmt='General"″"')
    r += 1
    merge(f"C{r}:J{r}", "Kemiringan pipa yang tersedia pada Tabel 16: 1 %, 2 %, dan 4 %. Pipa cabang dicek terhadap luas per roof "
                        "drain terbesar; pipa gabungan dicek terhadap luas atap total.", font=F_I, al=_LEFT, border=False)
    ws.row_dimensions[r].height = 26
    r += 2

    # ---- 5. catatan drainase atap
    sec(5, "Catatan untuk Sistem Drainase Atap")
    put(f"C{r}", "A. Roof Drain PVC Ø (inci)", al=_LEFT)
    put(f"D{r}", P["d_roof"], F_IN, fmt='General"″"')
    put(f"E{r}", "Jumlah titik")
    put(f"F{r}", f"=SUM(F{z0}:F{z1})", F_B, fmt="0")
    r += 1
    merge(f"C{r}:J{r}", f'="Jika "&TEXT({ATOT},"0.00")&" m² dibagi "&TEXT(F{r - 1},"0")&" roof drain: setiap cabang melayani "&TEXT({ADRAIN},"0.00")&" m² ("&TEXT({ADRAIN},"0.00")&IF({ADRAIN}<=G{rc}," < "," > ")&TEXT(G{rc},"0")&"), "&IF({ADRAIN}<=G{rc},"tergolong layak dan aman untuk pipa Ø"&E{rc}&"″.","TIDAK layak untuk pipa Ø"&E{rc}&"″.")',
          font=F_N, fill=_SOFT, al=_LEFT)
    ws.row_dimensions[r].height = 30
    r += 1
    put(f"C{r}", "B. Pipa Air Hujan Horizontal PVC Ø (inci)", al=_LEFT)
    put(f"D{r}", f"=E{rg}", fmt='General"″"')
    r += 2

    # ---- 6. pipa tegak
    sec(6, "Pipa Kolektor Air Hujan Vertikal — SNI 8153:2015 (Tabel 17)")
    put(f"C{r}", "C. Pipa Kolektor Vertikal PVC Ø (inci)", al=_LEFT)
    put(f"D{r}", P["d_tegak"], F_IN, fmt='General"″"')
    rt = r
    r += 1
    put(f"C{r}", "Luas atap yang dilayani (m²)", al=_LEFT)
    put(f"D{r}", f"={ATOT}", fmt="0.00")
    r += 1
    put(f"C{r}", "Debit rencana (L/dt)", al=_LEFT)
    put(f"D{r}", f"={QLS}", fmt="0.00")
    r += 1
    put(f"C{r}", "Kapasitas Tabel 17 (m²) — isi dari SNI", al=_LEFT)
    put(f"D{r}", P["cap_tegak"] or 0, F_IN, fmt="#,##0")
    rcap = r
    r += 1
    put(f"C{r}", "Pemakaian kapasitas", al=_LEFT)
    put(f"D{r}", f'=IF(D{rcap}>0,D{rt + 1}/D{rcap},"–")', fmt="0%")
    r += 1
    put(f"C{r}", "Status", F_B, al=_LEFT)
    put(f"D{r}", f'=IF(D{rcap}>0,IF(D{rt + 1}<=D{rcap},"MEMENUHI","TIDAK MEMENUHI"),"BELUM DICEK")', F_B)
    rstat = r
    r += 2

    # ---- 7. kesimpulan
    sec(7, "Kesimpulan")
    kes = (f'="Penentuan ukuran roof drain dan pipa air hujan dilakukan berdasarkan luas bidang atap efektif ("&TEXT({ATOT},"0.00")&" m²), '
           f'intensitas hujan rencana "&TEXT({RI},"0")&" mm/jam, dan ketentuan SNI 8153:2015. "&'
           f'IF(AND(I{rc}="MEMENUHI",I{rg}="MEMENUHI"),"Hasil evaluasi menunjukkan bahwa pipa cabang Ø"&E{rc}&"″ dan pipa horizontal gabungan Ø"&E{rg}&"″ mampu melayani luas bidang atap rencana.",'
           f'"Hasil evaluasi menunjukkan ada pipa yang belum memenuhi persyaratan kapasitas; lihat kolom Ø minimum.")&'
           f'IF(D{rstat}="BELUM DICEK"," Kapasitas pipa tegak Ø"&D{rt}&"″ perlu dicek pada Tabel 17 SNI 8153:2015.",'
           f'IF(D{rstat}="MEMENUHI"," Pipa tegak memenuhi kapasitas Tabel 17."," Pipa tegak belum memenuhi kapasitas Tabel 17."))')
    merge(f"C{r}:J{r}", kes, font=F_N, fill=PatternFill("solid", fgColor="F1FAF5"), al=_LEFT)
    ws.row_dimensions[r].height = 78
    r += 2
    merge(f"B{r}:J{r}", "Keterangan: sel biru = input (boleh diubah); sel hitam = rumus otomatis. " + data["catatan_standar"],
          font=F_I, al=_LEFT, border=False)
    ws.row_dimensions[r].height = 26

    # ---- format kondisional
    for rng, col in ((f"I{rc}:I{rg}", "I"), (f"D{rstat}", "D")):
        first = rng.split(":")[0]
        for teks, warna, fc in [("MEMENUHI", "DDF3E6", "1E8E5A"), ("TIDAK MEMENUHI", "FCE8E6", "C0392B"), ("BELUM DICEK", "FFF1CC", "C77700")]:
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{first}="{teks}"'],
                                          fill=PatternFill("solid", bgColor=warna, fgColor=warna), font=Font(bold=True, color=fc)))
    ws.conditional_formatting.add(f"H{rc}:H{rg}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1.25,
                                                             color="35B6C9", showValue=True))
    ws.conditional_formatting.add(f"C{u0}:G{u1}", FormulaRule(formula=[f"OR($C{u0}=$E${rc},$C{u0}=$E${rg})"],
                                  fill=PatternFill("solid", bgColor="FFF6DA", fgColor="FFF6DA")))

    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# -------------------------------------------------------------------- PDF
def _cari_soffice() -> str | None:
    p = shutil.which("soffice") or shutil.which("libreoffice")
    if p:
        return p
    for cand in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                 r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if os.path.exists(cand):
            return cand
    return None


def excel_ke_pdf(xlsx_bytes: bytes) -> bytes | None:
    """xlsx -> pdf via LibreOffice. None kalau LibreOffice tidak terpasang / gagal."""
    exe = _cari_soffice()
    if not exe:
        return None
    with tempfile.TemporaryDirectory() as d:
        src = os.path.join(d, "laporan.xlsx")
        with open(src, "wb") as f:
            f.write(xlsx_bytes)
        try:
            subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", d, src],
                           check=True, timeout=120, capture_output=True)
        except (subprocess.SubprocessError, OSError):
            return None
        out = os.path.join(d, "laporan.pdf")
        if not os.path.exists(out):
            return None
        with open(out, "rb") as f:
            return f.read()
