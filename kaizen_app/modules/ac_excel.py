"""Modul AC - ekspor Excel 2 sheet (AC STANDART & AC FULL, rumus hidup) dan PDF."""
from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter

from modules.ac_calc import MODE_FULL, MODE_STANDAR

_THIN = Side(style="thin", color="444444")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_GREY = PatternFill("solid", fgColor="F2F2F2")
_BLUE = PatternFill("solid", fgColor="E8EEF7")
_MID = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center")
F_N = Font(name="Arial", size=10)
F_B = Font(name="Arial", size=10, bold=True)
F_IN = Font(name="Arial", size=10, color="0000FF")      # biru = input
F_I = Font(name="Arial", size=9, italic=True)

_KOLOM = {
    MODE_STANDAR: dict(dasar="L", total="L", jenis="M", konv="N", n="O", jml="P", status="Q", rek="R"),
    MODE_FULL: dict(dasar="L", orang="M", akt="N", worg="O", qorang="P", lamp="Q", lw="R", qlamp="S",
                    alat="T", qalat="U", total="V", jenis="W", konv="X", n="Y", jml="Z", status="AA", rek="AB"),
}
_LEBAR = {
    MODE_STANDAR: {"A": 2, "B": 5, "C": 30, "J": 20, "K": 13, "L": 14, "M": 14, "N": 11, "O": 8,
                   "P": 10, "Q": 12, "R": 14},
    MODE_FULL: {"A": 2, "B": 5, "C": 30, "J": 26, "K": 13, "L": 13, "M": 13, "N": 24, "O": 10, "P": 12,
                "Q": 9, "R": 10, "S": 12, "T": 11, "U": 12, "V": 13, "W": 14, "X": 11, "Y": 8, "Z": 10,
                "AA": 12, "AB": 14},
}


def _sheet(ws, hasil: pd.DataFrame, data: dict, proyek: dict, mode: str, od: float, fl: float):
    full = mode == MODE_FULL
    C = _KOLOM[mode]
    last_col = C["rek"]
    last_idx = column_index_from_string(last_col)
    kv, acs, fp = data["konversi"], data["ac"], data["full"]
    w2b_txt = str(kv["watt_ke_btu"]).replace(".", ",")

    def put(ref, val, font=F_N, fill=None, al=_MID, fmt=None, border=True):
        c = ws[ref]
        c.value = val
        c.font, c.alignment = font, al
        if border:
            c.border = _BOX
        if fill:
            c.fill = fill
        if fmt:
            c.number_format = fmt

    def merge(rng, val, **kw):
        ws.merge_cells(rng)
        put(rng.split(":")[0], val, **kw)
        if kw.get("border", True):
            for row in ws[rng]:
                for c in row:
                    c.border = _BOX

    # ---- identitas proyek
    for r, (lab, key) in enumerate([("PEKERJAAN :", "pekerjaan"), ("LOKASI :", "lokasi"),
                                    ("TAHUN :", "tahun"), ("ITEM PEKERJAAN :", "item")], start=2):
        merge(f"B{r}:C{r}", lab, font=F_B, al=_LEFT, border=False)
        merge(f"D{r}:{last_col}{r}", f": {proyek.get(key, '')}", font=F_N, al=_LEFT, border=False)
    put("B6", f"METODE: {mode.upper()}", F_B, al=_LEFT, border=False)

    # ---- keterangan + rumus (kiri)
    put("B8", "Dimana :", F_B, al=_LEFT, border=False)
    e_max = max(data["orientasi"].values())
    leg = [("N =", "Kapasitas AC (BTU/h)"), ("L =", "panjang ruangan (feet)"),
           ("W =", "lebar ruangan (feet)"), ("H =", "tinggi ruangan (feet)"),
           ("E =", "Faktor Orientasi matahari")]
    leg += [("", f"{k} → {v}" + (" (paling panas)" if v == e_max else "")) for k, v in data["orientasi"].items()]
    leg.append(("I =", "Faktor Isolasi posisi ruangan"))
    leg += [("", f"I = {v} → {k}") for k, v in data["isolasi"].items()]
    leg.append(("", f"1 meter = {str(kv['meter_ke_feet']).replace('.', ',')} Feet"))
    leg.append(("", ""))
    leg.append(("Rumus", ""))
    p = kv["pembagi"]
    if full:
        leg += [("", f"Q dasar = (L × W × H × I × E) / {p}"),
                ("", f"Q orang = max(0; n orang − orang dasar) × W/orang × {w2b_txt}"),
                ("", f"Q lampu = n lampu × W/lampu × F lampu × {w2b_txt}"),
                ("", f"Q peralatan = W peralatan × {w2b_txt}"),
                ("", "Q total = Q dasar + Q orang + Q lampu + Q peralatan"),
                ("", "N (AC) = Q total / kapasitas AC terpilih")]
    else:
        leg += [("", f"Btu/h = (L × W × H × I × E) / {p}"),
                ("", "N (AC) = Btu/h / kapasitas AC terpilih")]
    for i, (a, b) in enumerate(leg):
        put(f"B{9 + i}", a, F_B if a == "Rumus" else F_N, al=_LEFT, border=False)
        put(f"C{9 + i}", b, F_N, al=_LEFT, border=False)
    kiri_akhir = 9 + len(leg) - 1

    # ---- tabel parameter (kanan, J..M)
    for col, t in zip("JKL", ["Konversi", "Nilai", "Satuan"]):
        put(f"{col}8", t, F_B, _GREY)
    konv = [("1 Meter", kv["meter_ke_feet"], "Feet", "0.00"),
            ("1 m²", kv["m2_ke_ft2"], "ft²", "0.0000"),
            ("1 Btu/h", kv["btu_ke_pk"], "PK", "0.000000000"),
            ("Pembagi rumus", kv["pembagi"], "-", "0")]
    if full:
        konv.append(("1 Watt", kv["watt_ke_btu"], "Btu/h", "0.000"))
    for i, (lab, val, sat, nf) in enumerate(konv):
        r = 9 + i
        put(f"J{r}", lab, F_N, al=_LEFT)
        put(f"K{r}", val, F_IN, fmt=nf)
        put(f"L{r}", sat, F_N)
    FT, DIV = "$K$9", "$K$12"
    W2B = "$K$13"

    r = 9 + len(konv)
    put(f"J{r}", "Jenis AC", F_B, _GREY)
    put(f"K{r}", "Btu/h", F_B, _GREY)
    put(f"L{r}", "Watt (SLD)", F_B, _GREY)
    ac0 = r + 1
    ac1 = ac0 + len(acs) - 1
    for i, a in enumerate(acs):
        put(f"J{ac0 + i}", a["label"], F_IN, al=_LEFT)
        put(f"K{ac0 + i}", a["btu"], F_IN, fmt="#,##0")
        put(f"L{ac0 + i}", a.get("watt"), F_IN, fmt="#,##0")
    r = ac1 + 1
    put(f"J{r}", "Toleransi N (AC)", F_N, al=_LEFT)
    put(f"K{r}", data.get("toleransi_n", 0), F_IN, fmt="0.00")
    put(f"L{r}", "batas 'Batas'", F_N)
    TOL = f"$K${r}"

    OD = FL = akt0 = akt1 = None
    if full:
        r += 1
        put(f"J{r}", "Orang dasar", F_N, al=_LEFT)
        put(f"K{r}", od, F_IN, fmt="0")
        put(f"L{r}", "orang", F_N)
        OD = f"$K${r}"
        r += 1
        put(f"J{r}", "Faktor lampu", F_N, al=_LEFT)
        put(f"K{r}", fl, F_IN, fmt="0.00")
        put(f"L{r}", "(LED = 1,0)", F_N)
        FL = f"$K${r}"
        r += 2
        for col, t in zip("JKLM", ["Aktivitas orang", "Total (W)", "Sensible (W)", "Latent (W)"]):
            put(f"{col}{r}", t, F_B, _GREY)
        akt0 = r + 1
        for i, a in enumerate(fp["aktivitas"]):
            put(f"J{akt0 + i}", a["label"], F_IN, al=_LEFT)
            put(f"K{akt0 + i}", a["total"], F_IN)
            put(f"L{akt0 + i}", a["sensible"], F_IN)
            put(f"M{akt0 + i}", a["latent"], F_IN)
        akt1 = akt0 + len(fp["aktivitas"]) - 1
        r = akt1 + 1
        put(f"J{r}", f"Sumber: {fp['sumber_aktivitas']}", F_I, al=_LEFT, border=False)

    ref_r = r + 2
    for i, t in enumerate(data.get("referensi", [])):
        put(f"J{ref_r + i}", t, F_I, al=_LEFT, border=False)
    kanan_akhir = ref_r + len(data.get("referensi", [])) - 1

    # ---- tabel utama
    ttl = max(kiri_akhir, kanan_akhir) + 2
    put(f"B{ttl}", f"Metode Estimasi Kebutuhan Pendinginan Ruangan — {mode}", F_B, al=_LEFT, border=False)
    h1, h2 = ttl + 1, ttl + 2
    merge(f"B{h1}:B{h2}", "NO", font=F_B, fill=_GREY)
    merge(f"C{h1}:C{h2}", "NAMA RUANGAN", font=F_B, fill=_GREY)
    merge(f"D{h1}:F{h1}", "Dimensi Ruangan (meter)", font=F_B, fill=_GREY)
    merge(f"G{h1}:I{h1}", "Dimensi Ruangan (Feet)", font=F_B, fill=_GREY)
    merge(f"J{h1}:K{h1}", "FF", font=F_B, fill=_GREY)
    for col, t in zip("DEFGHI", ["L", "W", "H", "L", "W", "H"]):
        put(f"{col}{h2}", t, F_B, _GREY)
    put(f"J{h2}", "I (Insulasi)", F_B, _GREY)
    put(f"K{h2}", "E (Kondisi)", F_B, _GREY)
    merge(f"L{h1}:L{h2}", "Beban Dasar (Btu/h)" if full else "Total Btu/h", font=F_B, fill=_GREY)
    if full:
        merge(f"M{h1}:P{h1}", "ORANG", font=F_B, fill=_BLUE)
        for col, t in zip("MNOP", ["Jumlah", "Aktivitas", "W/orang", "Btu/h"]):
            put(f"{col}{h2}", t, F_B, _BLUE)
        merge(f"Q{h1}:S{h1}", "LAMPU", font=F_B, fill=_BLUE)
        for col, t in zip("QRS", ["Jumlah", "W/lampu", "Btu/h"]):
            put(f"{col}{h2}", t, F_B, _BLUE)
        merge(f"T{h1}:U{h1}", "PERALATAN", font=F_B, fill=_BLUE)
        for col, t in zip("TU", ["Daya (W)", "Btu/h"]):
            put(f"{col}{h2}", t, F_B, _BLUE)
        merge(f"{C['total']}{h1}:{C['total']}{h2}", "Total Btu/h", font=F_B, fill=_GREY)
    akhir = [("jenis", "Jenis AC yang akan digunakan (PK)"), ("konv", "Konversi (Btu/H)"), ("n", "N (AC)"),
             ("jml", "Jumlah AC Dalam Ruangan"), ("status", "Status"), ("rek", "Rekomendasi AC")]
    for key, t in akhir:
        merge(f"{C[key]}{h1}:{C[key]}{h2}", t, font=F_B, fill=_GREY)
    ws.row_dimensions[h1].height = 32
    ws.row_dimensions[h2].height = 32

    def rek_formula(tot, jml):
        expr = f'"> "&$J${ac1}&" (tambah unit)"'
        for i in range(ac1, ac0 - 1, -1):
            expr = f"IF({tot}<=$K${i}*MAX(1,{jml}),$J${i},{expr})"
        return "=" + expr

    r = h2 + 1
    first = r
    for no, (lantai, grup) in enumerate(hasil.groupby("Lantai", sort=False), start=1):
        put(f"B{r}", no, F_B)
        put(f"C{r}", lantai, Font(name="Arial", size=10, bold=True, color="0000FF"), al=_LEFT)
        for ci in range(4, last_idx + 1):
            ws[f"{get_column_letter(ci)}{r}"].border = _BOX
        r += 1
        for _, x in grup.iterrows():
            put(f"B{r}", None)
            put(f"C{r}", x["Ruangan"], F_IN, al=_LEFT)
            put(f"D{r}", x["L_m"], F_IN, fmt="0.00")
            put(f"E{r}", x["W_m"], F_IN, fmt="0.00")
            put(f"F{r}", x["H_m"], F_IN, fmt="0.00")
            put(f"G{r}", f"=D{r}*{FT}", fmt="0.00")
            put(f"H{r}", f"=E{r}*{FT}", fmt="0.00")
            put(f"I{r}", f"=F{r}*{FT}", fmt="0.00")
            put(f"J{r}", x["I"], F_IN)
            put(f"K{r}", x["E"], F_IN)
            put(f"L{r}", f"=(G{r}*H{r}*I{r}*J{r}*K{r})/{DIV}", fmt="#,##0.00")
            if full:
                put(f"M{r}", x["Orang"], F_IN, fmt="0")
                put(f"N{r}", x["Aktivitas"], F_IN)
                put(f"O{r}", f"=VLOOKUP(N{r},$J${akt0}:$K${akt1},2,FALSE)", fmt="#,##0")
                put(f"P{r}", f"=MAX(0,M{r}-{OD})*O{r}*{W2B}", fmt="#,##0.00")
                put(f"Q{r}", x["N_lampu"], F_IN, fmt="0")
                put(f"R{r}", x["W_lampu"], F_IN, fmt="0")
                put(f"S{r}", f"=Q{r}*R{r}*{FL}*{W2B}", fmt="#,##0.00")
                put(f"T{r}", x["W_alat"], F_IN, fmt="#,##0")
                put(f"U{r}", f"=T{r}*{W2B}", fmt="#,##0.00")
                put(f"{C['total']}{r}", f"=L{r}+P{r}+S{r}+U{r}", F_B, fmt="#,##0.00")
            t, j = f"{C['total']}{r}", f"{C['jml']}{r}"
            put(f"{C['jenis']}{r}", x["Jenis AC"], F_IN)
            put(f"{C['konv']}{r}", f"=VLOOKUP({C['jenis']}{r},$J${ac0}:$K${ac1},2,FALSE)", fmt="#,##0")
            put(f"{C['n']}{r}", f"={t}/{C['konv']}{r}", fmt="0.000")
            put(j, int(x["Jumlah"]), F_IN)
            nn = f"{C['n']}{r}"
            put(f"{C['status']}{r}", f'=IF({j}<=0,"Isi jumlah AC",IF({nn}/{j}<=1,"Cukup",'
                                     f'IF({nn}/{j}<=1+{TOL},"Batas","Kurang")))', F_B)
            put(f"{C['rek']}{r}", rek_formula(t, j), F_B)
            r += 1
    last = r - 1

    # ---- total
    merge(f"B{r}:K{r}", "TOTAL", font=F_B, fill=_GREY)
    sums = ["L"] + (["M", "P", "Q", "S", "T", "U"] if full else [])
    sums += [C["total"], C["jml"]] if full else [C["jml"]]
    for ci in range(column_index_from_string("L"), last_idx + 1):
        col = get_column_letter(ci)
        if col in sums or (not full and col == C["total"]):
            put(f"{col}{r}", f"=SUM({col}{first}:{col}{last})", F_B, _GREY, fmt="#,##0.00" if col not in (C["jml"], "M", "Q") else "0")
        else:
            put(f"{col}{r}", None, F_B, _GREY)
    tot_r = r

    rng = f"{C['status']}{first}:{C['status']}{last}"
    for teks, warna, fc in [("Cukup", "E6F4EA", "137333"), ("Batas", "FEF7E0", "B06000"),
                            ("Kurang", "FCE8E6", "C5221F")]:
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'${C["status"]}{first}="{teks}"'],
            fill=PatternFill("solid", bgColor=warna, fgColor=warna), font=Font(bold=True, color=fc)))

    # ---- rekap jenis AC
    rk = tot_r + 3
    put(f"B{rk}", "REKAP JENIS AC (acuan untuk Modul Beban Listrik / SLD)", F_B, al=_LEFT, border=False)
    for col, t in zip("CDEF", ["Jenis AC", "Unit", "Daya/unit (W)", "Total (W)"]):
        put(f"{col}{rk + 1}", t, F_B, _GREY)
    jc, pc = C["jenis"], C["jml"]
    for i in range(len(acs)):
        x = rk + 2 + i
        a_row = ac0 + i
        put(f"C{x}", f"=J{a_row}", F_N, al=_LEFT)
        put(f"D{x}", f"=SUMIF(${jc}${first}:${jc}${last},C{x},${pc}${first}:${pc}${last})")
        put(f"E{x}", f'=IF(ISNUMBER(L{a_row}),L{a_row},"-")', fmt="#,##0")
        put(f"F{x}", f'=IF(ISNUMBER(E{x}),D{x}*E{x},"-")', fmt="#,##0")
    x0, x1 = rk + 2, rk + 1 + len(acs)
    put(f"C{x1 + 1}", "TOTAL", F_B, _GREY, al=_LEFT)
    put(f"D{x1 + 1}", f"=SUM(D{x0}:D{x1})", F_B, _GREY)
    put(f"E{x1 + 1}", None, F_B, _GREY)
    put(f"F{x1 + 1}", f"=SUM(F{x0}:F{x1})", F_B, _GREY, fmt="#,##0")

    n_r = x1 + 3
    if data.get("catatan"):
        merge(f"B{n_r}:{last_col}{n_r}", f"“{data['catatan']}”", font=F_I, al=_LEFT, border=False)
    put(f"B{n_r + 1}", "Keterangan: sel biru = input (boleh diubah); sel hitam = rumus otomatis. "
                       "Status 'Batas' = N/jumlah AC melebihi 1 tetapi masih dalam toleransi.",
        F_I, al=_LEFT, border=False)
    if full:
        put(f"B{n_r + 2}", "Catatan Full Calculation: beban selubung rinci (CLTD) dan udara segar ventilasi "
                           "belum dihitung; orang dasar dianggap sudah termasuk pada rumus dasar.",
            F_I, al=_LEFT, border=False)

    # ---- lebar kolom & cetak
    for ci in range(1, last_idx + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 9 if ci > 3 else 5
    for col, w in _LEBAR[mode].items():
        ws.column_dimensions[col].width = w
    # kolom Dimensi Feet (G:I) disembunyikan: tetap dipakai rumus, tidak tampil/tercetak
    for col in "GHI":
        ws.column_dimensions[col].hidden = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3 if full else ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1 if (n_r + 3) <= 75 else 0   # sheet pendek -> 1 halaman
    ws.print_options.horizontalCentered = True


def build_excel(hasil_standar: pd.DataFrame, hasil_full: pd.DataFrame, data: dict, proyek: dict,
                orang_dasar: float | None = None, faktor_lampu: float | None = None) -> bytes:
    od = data["full"]["orang_dasar"] if orang_dasar is None else orang_dasar
    fl = data["full"]["faktor_lampu"] if faktor_lampu is None else faktor_lampu
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "AC STANDART"
    _sheet(ws1, hasil_standar, data, proyek, MODE_STANDAR, od, fl)
    ws2 = wb.create_sheet("AC FULL")
    _sheet(ws2, hasil_full, data, proyek, MODE_FULL, od, fl)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# -------------------------------------------------------------------- PDF
def _cari_soffice() -> str | None:
    p = shutil.which("soffice") or shutil.which("libreoffice")
    if p:
        return p
    for cand in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                 r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if os.path.exists(cand):
            return cand
    return None


def excel_ke_pdf(xlsx_bytes: bytes) -> bytes | None:
    """xlsx -> pdf via LibreOffice. None kalau LibreOffice tidak terpasang / gagal."""
    exe = _cari_soffice()
    if not exe:
        return None
    with tempfile.TemporaryDirectory() as d:
        src = os.path.join(d, "laporan.xlsx")
        with open(src, "wb") as f:
            f.write(xlsx_bytes)
        try:
            subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", d, src],
                           check=True, timeout=120, capture_output=True)
        except (subprocess.SubprocessError, OSError):
            return None
        out = os.path.join(d, "laporan.pdf")
        if not os.path.exists(out):
            return None
        with open(out, "rb") as f:
            return f.read()
