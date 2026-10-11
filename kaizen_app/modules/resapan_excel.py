"""Modul Resapan - ekspor Excel (rumus hidup, desain konsisten) dan PDF."""
from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from io import BytesIO

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

from modules.resapan_calc import BELUM, OK, TIDAK
from modules.resapan_skema import skema_png

NAVY, AQUA, SOFT, LINE = "0F3D63", "35B6C9", "EAF3F9", "D5E1EB"
_THIN = Side(style="thin", color=LINE)
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_F_FILL = PatternFill("solid", fgColor=NAVY)
_SOFT = PatternFill("solid", fgColor=SOFT)
_MID = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_NW = Alignment(horizontal="left", vertical="center", wrap_text=False)
F_N = Font(name="Arial", size=10)
F_B = Font(name="Arial", size=10, bold=True)
F_IN = Font(name="Arial", size=10, color="0000FF")      # biru = input
F_H = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F_I = Font(name="Arial", size=9, italic=True, color="6B7C8C")
F_T = Font(name="Arial", size=15, bold=True, color="FFFFFF")
F_S = Font(name="Arial", size=12, bold=True, color=NAVY)
REF = "Referensi"


def _put(ws, ref, val, font=F_N, fill=None, al=_MID, fmt=None, border=True):
    c = ws[ref]
    c.value = val
    c.font, c.alignment = font, al
    if border:
        c.border = _BOX
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt


def _merge(ws, rng, val, **kw):
    ws.merge_cells(rng)
    _put(ws, rng.split(":")[0], val, **kw)
    if kw.get("border", True):
        for row in ws[rng]:
            for c in row:
                c.border = _BOX


def _judul(ws, rng, teks, tinggi=32):
    _merge(ws, rng, teks, font=F_T, fill=_F_FILL, al=_NW, border=False)
    ws.row_dimensions[int(rng.split(":")[0][1:])].height = tinggi


def _cf_status(ws, rng, first):
    for teks, warna, fc in [(OK, "DDF3E6", "1E8E5A"), (TIDAK, "FCE8E6", "C0392B"), (BELUM, "FFF1CC", "C77700"),
                            ("DIHITUNG", "DDF3E6", "1E8E5A"), ("TIDAK DIHITUNG", "ECEFF2", "6B7C8C")]:
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{first}="{teks}"'],
                                      fill=PatternFill("solid", bgColor=warna, fgColor=warna), font=Font(bold=True, color=fc)))


def build_excel(h: dict, data: dict, proyek: dict) -> bytes:
    P, v, sni = h["P"], h["vab"], data["sni"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Resapan"
    wk = wb.create_sheet("Skema")
    wr = wb.create_sheet(REF)
    wn = wb.create_sheet("SNI & Peraturan")
    put = lambda *a, **k: _put(ws, *a, **k)
    merge = lambda *a, **k: _merge(ws, *a, **k)
    yn = lambda b: "Ya" if b else "Tidak"
    dv = DataValidation(type="list", formula1='"Ya,Tidak"', allow_blank=False)
    ws.add_data_validation(dv)

    # ===================================== Referensi
    _judul(wr, "B1:E1", "REFERENSI: KOEFISIEN LIMPASAN DAN JARAK MINIMUM")
    for col, tx in zip("BCD", ["Permukaan", "C", "Bidang tadah (1 = ya)"]):
        _put(wr, f"{col}3", tx, F_H, _F_FILL)
    for i, p in enumerate(data["permukaan"]):
        r = 4 + i
        _put(wr, f"B{r}", p["nama"], F_B, al=_LEFT)
        _put(wr, f"C{r}", p["c"], F_IN, fmt="0.00")
        _put(wr, f"D{r}", 1 if p["tadah"] else 0, F_IN, fmt="0")
    rend = 3 + len(data["permukaan"])
    _merge(wr, f"B{rend + 1}:E{rend + 1}", data["catatan_c"], font=F_I, al=_LEFT, border=False)
    wr.row_dimensions[rend + 1].height = 30
    r0 = rend + 3
    for col, tx in zip("BCD", ["Jarak minimum sumur resapan terhadap", "Jarak (m)", "Acuan"]):
        _put(wr, f"{col}{r0}", tx, F_H, _F_FILL)
    for i, j in enumerate(sni["jarak"]):
        _put(wr, f"B{r0 + 1 + i}", j["nama"], F_N, al=_LEFT)
        _put(wr, f"C{r0 + 1 + i}", j["m"], F_IN, fmt="0")
        _put(wr, f"D{r0 + 1 + i}", "SNI 03-2453-2002 Tabel 1")
    _merge(wr, f"B{r0 + 5}:E{r0 + 5}", data["hujan"]["catatan"], font=F_I, al=_LEFT, border=False)
    wr.row_dimensions[r0 + 5].height = 52
    for col, w in {"A": 2, "B": 48, "C": 12, "D": 24, "E": 20}.items():
        wr.column_dimensions[col].width = w
    TAB = f"{REF}!$B$4:$D${rend}"

    # ===================================== SNI & peraturan
    _judul(wn, "B1:F1", "LANDASAN PERENCANAAN: SNI & PERATURAN TERKAIT RESAPAN AIR HUJAN")
    for col, tx in zip("BCDEF", ["No", "Standar / Peraturan", "Judul", "Peran pada perhitungan", "Catatan"]):
        _put(wn, f"{col}3", tx, F_H, _F_FILL)
    for i, s in enumerate(data["standar"]):
        r = 4 + i
        _put(wn, f"B{r}", i + 1)
        _put(wn, f"C{r}", s["kode"], F_B, al=_LEFT)
        _put(wn, f"D{r}", s["judul"], F_N, al=_LEFT)
        _put(wn, f"E{r}", s["peran"], F_N, al=_LEFT)
        _put(wn, f"F{r}", s["catatan"], F_N, al=_LEFT)
        wn.row_dimensions[r].height = 58
    _put(wn, f"B{5 + len(data['standar'])}", data["catatan_standar"], F_I, al=_NW, border=False)
    for col, w in {"A": 2, "B": 5, "C": 30, "D": 44, "E": 56, "F": 50}.items():
        wn.column_dimensions[col].width = w
    wn.page_setup.orientation = "landscape"
    wn.sheet_properties.pageSetUpPr.fitToPage = True
    wn.page_setup.fitToWidth, wn.page_setup.fitToHeight = 1, 0

    # ===================================== Skema
    _judul(wk, "B1:N1", "SKEMA OPSI RESAPAN (ILUSTRASI PROPORSIONAL, BUKAN GAMBAR KERJA)")
    img = XLImage(BytesIO(skema_png(h)))
    ratio = img.height / max(1, img.width)
    img.width, img.height = 1150, int(1150 * ratio)
    wk.add_image(img, "B3")
    wk.sheet_view.showGridLines = False
    wk.page_setup.orientation = "landscape"
    wk.sheet_properties.pageSetUpPr.fitToPage = True
    wk.page_setup.fitToWidth, wk.page_setup.fitToHeight = 1, 1

    # ===================================== sheet utama
    for col, w in {"A": 2, "B": 5, "C": 50, "D": 16, "E": 17, "F": 14, "G": 50}.items():
        ws.column_dimensions[col].width = w
    _judul(ws, "B1:G1", "PERHITUNGAN PENGELOLAAN AIR HUJAN — SUMUR RESAPAN DAN BIOPORI", 34)
    for r, (lab, key) in enumerate([("PEKERJAAN", "pekerjaan"), ("LOKASI", "lokasi"), ("TAHUN", "tahun"),
                                    ("ITEM PEKERJAAN", "item")], start=3):
        merge(f"B{r}:C{r}", lab, font=Font(name="Arial", size=10, bold=True, color=NAVY), fill=_SOFT, al=_NW, border=False)
        merge(f"D{r}:G{r}", f": {proyek.get(key, '')}", font=F_N, fill=_SOFT, al=_NW, border=False)
    merge("B8:G8", "Landasan — SNI 03-2453-2002 · Permen PU 11/PRT/M/2014 · Permen LH 12/2009 (biopori) · "
                   "lengkap pada sheet 'SNI & Peraturan'", font=F_I, al=_LEFT, border=False)
    r = 10

    def sec(no, judul):
        nonlocal r
        put(f"B{r}", no, F_H, _F_FILL)
        merge(f"C{r}:G{r}", judul, font=F_S, al=_NW, border=False)
        for col in "CDEFG":
            ws[f"{col}{r}"].border = Border(bottom=Side(style="medium", color=AQUA))
        ws.row_dimensions[r].height = 22
        r += 1

    def baris(label, val, unit, rumus="", font=F_N, fmt="#,##0.000", fill=None, inp=False, validasi=False):
        nonlocal r
        put(f"C{r}", label, font if not inp else F_N, fill, al=_LEFT)
        put(f"D{r}", val, F_IN if inp else font, fill, fmt=fmt)
        put(f"E{r}", unit, F_N, fill)
        merge(f"F{r}:G{r}", rumus, font=F_I if not fill else F_B, fill=fill, al=_LEFT)
        if validasi:
            dv.add(ws[f"D{r}"])
        ref = f"$D${r}"
        r += 1
        return ref

    # ---- 1. Vab
    sec(1, "Volume Andil Banjir / Volume Wajib Kelola — SNI 03-2453-2002 (1), Permen PU 11/PRT/M/2014")
    merge(f"C{r}:G{r}", "Vab = 0,855 × Σ (C × A) × R / 1000", font=Font(name="Cambria", size=13, italic=True, color=NAVY),
          fill=_SOFT, al=_MID)
    ws.row_dimensions[r].height = 26
    r += 1
    FV = baris("Faktor Vab", sni["faktor_vab"], "–", "konstanta rumus SNI 03-2453-2002", fmt="0.000", inp=True)
    RR = baris("Tinggi hujan harian rata-rata (R)", P["R"], "mm/hari", f"{data['hujan']['kota']} (Excel); = L/m²/hari", fmt="0.0", inp=True)
    TM = baris("Taman dihitung dalam Vab?", yn(P["taman"]), "Ya/Tidak", "Tidak bila pekarangan hijau menyerap air", fmt="@", inp=True, validasi=True)
    r += 1
    for col, tx in zip("BCDEFG", ["NO", "PERMUKAAN", "Luas A (m²)", "C", "C × A", "Dalam Vab"]):
        put(f"{col}{r}", tx, F_H, _F_FILL)
    r += 1
    s0 = r
    for i, x in enumerate(v["rinc"].itertuples(), start=1):
        put(f"B{r}", i)
        put(f"C{r}", x.Permukaan, F_IN, al=_LEFT)
        put(f"D{r}", x.A, F_IN, fmt="#,##0.00")
        put(f"E{r}", f"=VLOOKUP(C{r},{TAB},2,FALSE)", fmt="0.00")
        put(f"G{r}", f'=IF(OR(VLOOKUP(C{r},{TAB},3,FALSE)=1,{TM}="Ya"),"DIHITUNG","TIDAK DIHITUNG")', F_B)
        put(f"F{r}", f'=IF(G{r}="DIHITUNG",D{r}*E{r},0)', fmt="0.000")
        r += 1
    s1 = r - 1
    _cf_status(ws, f"G{s0}:G{s1}", f"G{s0}")
    merge(f"B{r}:C{r}", "TOTAL BIDANG TADAH", font=F_B, fill=_SOFT, al=_NW)
    put(f"D{r}", f'=SUMIFS(D{s0}:D{s1},G{s0}:G{s1},"DIHITUNG")', F_B, _SOFT, fmt="#,##0.00")
    put(f"E{r}", f'=IF(D{r}>0,F{r}/D{r},"")', F_B, _SOFT, fmt="0.00")
    put(f"F{r}", f"=SUM(F{s0}:F{s1})", F_B, _SOFT, fmt="0.000")
    put(f"G{r}", "C rata-rata tertimbang", F_I, _SOFT)
    ATD, CA = f"$D${r}", f"$F${r}"
    r += 2
    VAB = baris("VOLUME ANDIL BANJIR (Vab)", f"={FV}*{CA}*{RR}/1000", "m³", "0,855 × Σ(C×A) × R / 1000", F_B, "0.000", _SOFT)
    baris("Vab", f"={VAB}*1000", "liter", "", fmt="#,##0")
    ROT = baris("Pembanding aturan praktis (bukan SNI)", f"={ATD}/{sni['rot_m2_per_m3']}", "m³",
                f"1 m³ resapan per {sni['rot_m2_per_m3']} m² bidang tadah (Excel menulis m³, seharusnya m²)", fmt="0.000")
    r += 1

    # ---- 2. tanah & persyaratan
    sec(2, "Data Tanah dan Persyaratan Teknis — SNI 03-2453-2002 pasal 4")
    KC = baris("Permeabilitas tanah (K)", P["k_cmjam"], "cm/jam", "hasil uji lapangan; syarat ≥ 2,0 cm/jam", fmt="0.00", inp=True)
    KV = baris("Kv (alas sumur)", f"={KC}*24/100", "m/hari", "K × 24 / 100", fmt="0.000")
    RK = baris("Rasio Kh / Kv (dinding terhadap alas)", P["rasio_kh"], "–", "contoh SNI: Kh = 2 × Kv", fmt="0.0", inp=True)
    KH = baris("Kh (dinding sumur)", f"={RK}*{KV}", "m/hari", "rasio × Kv", fmt="0.000")
    MT = baris("Kedalaman muka air tanah (MAT)", P["mat"], "m", "isi dari pengukuran; 0 = belum diisi (syarat ≥ 1,5 m)", fmt="0.00", inp=True)
    TE = baris("Durasi hujan efektif (te)", f"={sni['te_koef']}*{RR}^{sni['te_pangkat']}/60", "jam", "te = 0,9 × R^0,92 / 60", fmt="0.000")
    r += 1
    for col, tx in zip("CDEFG", ["Persyaratan", "Isian", "Status", "Acuan", ""]):
        if tx or col == "G":
            put(f"{col}{r}", tx, F_H, _F_FILL)
    r += 1
    q0 = r
    cek = P["cek"]
    hmax = f"MAX({{H1}},{{H2}})"
    daftar = [
        (f"Permeabilitas tanah ≥ {sni['k_min_cm_jam']:g} cm/jam", None, f'=IF({KC}>={sni["k_min_cm_jam"]},"{OK}","{TIDAK}")', "SNI 03-2453-2002 pasal 4.2"),
        (f"Muka air tanah ≥ {sni['mat_min']:g} m (musim hujan)", None, f'=IF({MT}>0,IF({MT}>={sni["mat_min"]},"{OK}","{TIDAK}"),"{BELUM}")', "SNI 03-2453-2002 pasal 4.2"),
        ("Kedalaman sumur rencana < muka air tanah", None, "{HMAT}", "SNI 03-2453-2002 pasal 5.2"),
        ("Lahan relatif datar", yn(cek["datar"]), "man", "SNI 03-2453-2002 pasal 4.1"),
        ("Air yang masuk adalah air hujan tidak tercemar", yn(cek["tak_tercemar"]), "man", "SNI 03-2453-2002 pasal 4.1"),
        (f"Jarak ke pondasi bangunan ≥ {sni['jarak'][1]['m']} m", yn(cek["pondasi"]), "man", "SNI 03-2453-2002 Tabel 1"),
        (f"Jarak ke sumur air bersih / sumur resapan lain ≥ {sni['jarak'][0]['m']} m", yn(cek["sumur_air"]), "man", "SNI 03-2453-2002 Tabel 1"),
        (f"Jarak ke bidang/sumur resapan tangki septik ≥ {sni['jarak'][2]['m']} m", yn(cek["septik"]), "man", "SNI 03-2453-2002 Tabel 1"),
        ("Memenuhi peraturan daerah setempat", yn(cek["perda"]), "man", "SNI 03-2453-2002 pasal 4.1"),
    ]
    q_row_hmat = None
    for teks, isian, fx, acuan in daftar:
        put(f"C{r}", teks, al=_LEFT)
        if isian is not None:
            put(f"D{r}", isian, F_IN)
            dv.add(ws[f"D{r}"])
            fx = f'=IF(D{r}="Ya","{OK}","{BELUM}")'
        else:
            put(f"D{r}", "otomatis", F_I)
        put(f"E{r}", fx if fx != "{HMAT}" else None, F_B)
        if fx == "{HMAT}":
            q_row_hmat = r
        merge(f"F{r}:G{r}", acuan, font=F_I, al=_LEFT)
        r += 1
    q1 = r - 1
    _cf_status(ws, f"E{q0}:E{q1}", f"E{q0}")
    r += 1

    # ---- sumur (lingkaran & persegi)
    def blok_sumur(no, judul, bentuk):
        nonlocal r
        sec(no, judul)
        merge(f"C{r}:G{r}", "N = H total / H rencana,   H total = (Vab − Vrsp) / A alas,   Vrsp = te/24 × A total × K rata-rata",
              font=Font(name="Cambria", size=12, italic=True, color=NAVY), fill=_SOFT, al=_MID)
        ws.row_dimensions[r].height = 24
        r += 1
        if bentuk == "lingkaran":
            D1 = baris("Diameter sumur (D)", P["D"], "m", "buis beton / pasangan", fmt="0.00", inp=True)
            H1 = baris("Kedalaman rencana (H)", P["H1"], "m", "< kedalaman muka air tanah", fmt="0.00", inp=True)
            KD = baris("Dinding kedap?", yn(P["kedap"]), "Ya/Tidak", "Ya = hanya alas yang meresap", fmt="@", inp=True, validasi=True)
            AA = baris("Luas alas (A alas)", f"=PI()*{D1}^2/4", "m²", "¼ π D²", fmt="0.000")
            AD = baris("Luas dinding (A dinding)", f"=PI()*{D1}*{H1}", "m²", "π D H", fmt="0.000")
        else:
            P1 = baris("Panjang sumur (P)", P["P"], "m", "", fmt="0.00", inp=True)
            L1 = baris("Lebar sumur (L)", P["L"], "m", "", fmt="0.00", inp=True)
            H1 = baris("Kedalaman rencana (H)", P["H2"], "m", "< kedalaman muka air tanah", fmt="0.00", inp=True)
            KD = baris("Dinding kedap?", yn(P["kedap"]), "Ya/Tidak", "Ya = hanya alas yang meresap", fmt="@", inp=True, validasi=True)
            AA = baris("Luas alas (A alas)", f"={P1}*{L1}", "m²", "P × L", fmt="0.000")
            AD = baris("Luas dinding (A dinding)", f"=2*({P1}+{L1})*{H1}", "m²", "2 (P + L) H", fmt="0.000")
        AT = baris("Luas resap total (A total)", f'=IF({KD}="Ya",{AA},{AA}+{AD})', "m²", "A alas + A dinding (alas saja bila dinding kedap)", fmt="0.000")
        KR = baris("K rata-rata", f'=IF({KD}="Ya",{KV},({KV}*{AA}+{KH}*{AD})/{AT})', "m/hari", "(Kv·A alas + Kh·A dinding) / A total   (3)", fmt="0.000")
        VR = baris("Volume meresap (Vrsp)", f"={TE}/24*{AT}*{KR}", "m³", "te/24 × A total × K rata-rata   (2)", fmt="0.000")
        VS = baris("Volume storasi", f"={VAB}-{VR}", "m³", "Vab − Vrsp   (4)", fmt="0.000")
        HT = baris("Kedalaman total (H total)", f"={VS}/{AA}", "m", "V storasi / A alas   (5)", fmt="0.00")
        NR = baris("n = H total / H rencana", f"={HT}/{H1}", "buah", "(6)", fmt="0.00")
        N = baris("JUMLAH SUMUR (dibulatkan ke atas)", f"=IF({NR}>0,MAX(1,ROUNDUP(ROUND({NR},9),0)),1)", "buah", "N = ⌈n⌉", F_B, "0", _SOFT)
        VU = baris("Volume tampung / sumur", f"={AA}*{H1}", "m³", "A alas × H", fmt="0.000")
        KP = baris("Kapasitas total", f"={N}*({VU}+{VR})", "m³", "N × (volume tampung + Vrsp)   — kontrol tambahan", fmt="0.000")
        RS = baris("Kapasitas / Vab", f"={KP}/{VAB}", "–", "≥ 100% = memenuhi", fmt="0%")
        ws.conditional_formatting.add(RS.replace("$", ""), DataBarRule(start_type="num", start_value=0, end_type="num", end_value=2,
                                                                      color="35B6C9", showValue=True))
        r += 1
        return {"N": N, "KP": KP, "RS": RS, "H": H1, "VU": VU}

    ws.row_breaks.append(Break(id=r - 1))
    L = blok_sumur(3, "Sumur Resapan Berpenampang Lingkaran", "lingkaran")
    Q = blok_sumur(4, "Sumur Resapan Berpenampang Persegi", "persegi")

    # status kedalaman < MAT (butuh H dari kedua blok)
    if q_row_hmat:
        put(f"E{q_row_hmat}", f'=IF({MT}>0,IF(MAX({L["H"]},{Q["H"]})<{MT},"{OK}","{TIDAK}"),"{BELUM}")', F_B)

    # ---- biopori
    ws.row_breaks.append(Break(id=r - 1))
    sec(5, "Resapan Biopori (Permen LH 12/2009, Brata & Nelistya 2008)")
    DB = baris("Diameter pipa biopori (d)", P["d_bio"], "m", "acuan 10–25 cm; pipa 4\" = 0,1016 m", fmt="0.0000", inp=True)
    TB = baris("Kedalaman (t)", P["t_bio"], "m", "acuan ± 100 cm; periksa muka air tanah", fmt="0.00", inp=True)
    VL = baris("Volume satu lubang", f"=PI()*({DB}/2)^2*{TB}", "m³", "π r² t", fmt="0.00000")
    baris("Volume satu lubang", f"={VL}*1000", "liter", "", fmt="0.00")
    ABA = f"PI()*{DB}^2/4"
    ABD = f"PI()*{DB}*{TB}"
    KRB = baris("K rata-rata", f"=({KV}*{ABA}+{KH}*{ABD})/({ABA}+{ABD})", "m/hari", "adaptasi SNI 03-2453-2002 pada lubang kecil", fmt="0.000")
    VRB = baris("Volume meresap / lubang", f"={TE}/24*({ABA}+{ABD})*{KRB}", "m³", "te/24 × A total × K rata-rata", fmt="0.00000")
    r += 1
    N1 = baris("A. N metode volume (konservatif, seperti Excel)", f"=ROUNDUP(ROUND({VAB}/{VL},9),0)", "lubang", "N = Vab / V lubang", F_B, "0", _SOFT)
    N2 = baris("B. N dengan peresapan (adaptasi SNI)", f"=ROUNDUP(ROUND({VAB}/({VL}+{VRB}),9),0)", "lubang", "N = Vab / (V lubang + Vrsp)", F_B, "0", _SOFT)
    II = baris("Intensitas hujan jam-jaman (I) — opsional", P["I_jam"], "mm/jam", "untuk metode C; isi bila ada", fmt="0.0", inp=True)
    PL = baris("Laju peresapan per lubang (P) — opsional", P["P_lph"], "liter/jam", "diukur di lapangan", fmt="0.0", inp=True)
    N3 = baris("C. N Brata & Nelistya", f'=IF(AND({II}>0,{PL}>0),ROUNDUP(ROUND({II}*{ATD}/{PL},9),0),"isi I dan P")', "lubang", "N = I × A / P", F_B, "0", _SOFT)
    NRT = baris("Pembanding aturan praktis (bukan SNI)", f"=ROUNDUP(ROUND({ROT}/{VL},9),0)", "lubang", "(A/25) / V lubang", fmt="0")
    HG = baris("Harga satuan pipa biopori", P["harga"], "Rp", "dari Excel; ganti sesuai harga proyek", fmt="#,##0", inp=True)
    baris("Estimasi biaya metode A", f"={N1}*{HG}", "Rp", "N × harga satuan (memakai N dibulatkan ke atas)", fmt="#,##0")
    baris("Estimasi biaya metode B", f"={N2}*{HG}", "Rp", "", fmt="#,##0")
    r += 1

    # ---- perbandingan
    sec(6, "Perbandingan Opsi Resapan")
    for col, tx in zip("CDEFG", ["Opsi", "Jumlah unit", "Volume tampung (m³)", "Kapasitas total (m³)", "Kapasitas / Vab"]):
        put(f"{col}{r}", tx, F_H, _F_FILL)
    ws.row_dimensions[r].height = 30
    r += 1
    c0 = r
    for nama, n, vt, kp in [("Sumur resapan lingkaran", L["N"], f"={L['N']}*{L['VU']}", f"={L['KP']}"),
                            ("Sumur resapan persegi", Q["N"], f"={Q['N']}*{Q['VU']}", f"={Q['KP']}"),
                            ("Biopori — metode volume", N1, f"={N1}*{VL}", f"={N1}*{VL}"),
                            ("Biopori — dengan peresapan", N2, f"={N2}*{VL}", f"={N2}*({VL}+{VRB})")]:
        put(f"C{r}", nama, F_B, al=_LEFT)
        put(f"D{r}", f"={n}", F_B, fmt="0")
        put(f"E{r}", vt, fmt="0.00")
        put(f"F{r}", kp, fmt="0.00")
        put(f"G{r}", f"=F{r}/{VAB}", fmt="0%")
        r += 1
    ws.conditional_formatting.add(f"G{c0}:G{r - 1}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=2,
                                                                color="35B6C9", showValue=True))
    r += 1

    sec(7, "Kesimpulan")
    kes = (f'="Volume andil banjir (Vab) = "&TEXT({VAB},"0.000")&" m³ (luas tadah "&TEXT({ATD},"0.00")&" m², R = "&TEXT({RR},"0")&" mm/hari). '
           f'Kebutuhan resapan: "&{L["N"]}&" sumur lingkaran, atau "&{Q["N"]}&" sumur persegi, atau "&{N1}&" lubang biopori (metode volume; "&{N2}&" lubang bila memperhitungkan peresapan). '
           f'Pembanding aturan praktis bukan bagian SNI."')
    merge(f"C{r}:G{r}", kes, font=F_N, fill=PatternFill("solid", fgColor="F1FAF5"), al=_LEFT)
    ws.row_dimensions[r].height = 62
    r += 2
    merge(f"B{r}:G{r}", "Keterangan: sel biru = input (boleh diubah); sel hitam = rumus otomatis. " + data["catatan_standar"],
          font=F_I, al=_LEFT, border=False)
    ws.row_dimensions[r].height = 28

    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_options.horizontalCentered = True

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# -------------------------------------------------------------------- PDF
def _cari_soffice() -> str | None:
    p = shutil.which("soffice") or shutil.which("libreoffice")
    if p:
        return p
    for cand in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                 r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if os.path.exists(cand):
            return cand
    return None


def excel_ke_pdf(xlsx_bytes: bytes) -> bytes | None:
    """xlsx -> pdf via LibreOffice. None kalau LibreOffice tidak terpasang / gagal."""
    exe = _cari_soffice()
    if not exe:
        return None
    with tempfile.TemporaryDirectory() as d:
        src = os.path.join(d, "laporan.xlsx")
        with open(src, "wb") as f:
            f.write(xlsx_bytes)
        try:
            subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", d, src],
                           check=True, timeout=120, capture_output=True)
        except (subprocess.SubprocessError, OSError):
            return None
        out = os.path.join(d, "laporan.pdf")
        if not os.path.exists(out):
            return None
        with open(out, "rb") as f:
            return f.read()
