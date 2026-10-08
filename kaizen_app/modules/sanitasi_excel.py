"""Modul Sanitasi - ekspor Excel (rumus hidup, mengikuti sheet 'Sanitasi+Air') dan PDF."""
from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

_THIN = Side(style="thin", color="444444")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_GREY = PatternFill("solid", fgColor="F2F2F2")
_MID = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_NW = Alignment(horizontal="left", vertical="center", wrap_text=False)   # judul/catatan (tanpa wrap)
F_N = Font(name="Arial", size=10)
F_B = Font(name="Arial", size=10, bold=True)
F_IN = Font(name="Arial", size=10, color="0000FF")      # biru = input
F_I = Font(name="Arial", size=9, italic=True)
F_T = Font(name="Arial", size=12, bold=True)


def _put(ws, ref, val, font=F_N, fill=None, al=_MID, fmt=None, border=True):
    c = ws[ref]
    c.value = val
    c.font, c.alignment = font, al
    if border:
        c.border = _BOX
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt


def _merge(ws, rng, val, **kw):
    ws.merge_cells(rng)
    _put(ws, rng.split(":")[0], val, **kw)
    if kw.get("border", True):
        for row in ws[rng]:
            for c in row:
                c.border = _BOX


def build_excel(h: dict, data: dict, proyek: dict) -> bytes:
    P = h["P"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Sanitasi+Air"
    wt = wb.create_sheet("Tabel Pemakaian Air")
    wn = wb.create_sheet("SNI & Peraturan")
    put = lambda *a, **k: _put(ws, *a, **k)
    merge = lambda *a, **k: _merge(ws, *a, **k)

    # ------------------------------------------------ sheet tabel pemakaian air
    _put(wt, "B2", "SNI 03-7065-2005 Tabel 1 — Pemakaian air dingin minimum sesuai penggunaan gedung", F_T,
         al=_NW, border=False)
    for col, t in zip("BCDEF", ["No", "Penggunaan gedung", "Pemakaian air", "Satuan", "Satuan jumlah"]):
        _put(wt, f"{col}4", t, F_B, _GREY)
    for i, a in enumerate(data["pemakaian_air"]):
        r = 5 + i
        _put(wt, f"B{r}", i + 1)
        _put(wt, f"C{r}", a["fungsi"], F_N, al=_LEFT)
        _put(wt, f"D{r}", a["nilai"], F_IN, fmt="#,##0")
        _put(wt, f"E{r}", a["satuan"], F_N, al=_LEFT)
        _put(wt, f"F{r}", a["per"], F_N, al=_LEFT)
    t_end = 4 + len(data["pemakaian_air"])
    _put(wt, f"B{t_end + 2}", data["sumber_tabel"], F_I, al=_NW, border=False)
    _put(wt, f"B{t_end + 3}", data["catatan_tabel"], F_I, al=_NW, border=False)
    for col, w in {"A": 2, "B": 5, "C": 32, "D": 14, "E": 38, "F": 20}.items():
        wt.column_dimensions[col].width = w
    wt.page_setup.orientation = "landscape"
    wt.sheet_properties.pageSetUpPr.fitToPage = True
    wt.page_setup.fitToWidth, wt.page_setup.fitToHeight = 1, 0
    TAB = f"'Tabel Pemakaian Air'!$C$5:$F${t_end}"

    # ------------------------------------------------ sheet SNI & peraturan
    _put(wn, "B2", "LANDASAN PERENCANAAN: SNI & PERATURAN TERKAIT SANITASI", F_T, al=_NW, border=False)
    for col, t in zip("BCDEF", ["No", "Standar / Peraturan", "Judul", "Peran pada perhitungan", "Catatan"]):
        _put(wn, f"{col}4", t, F_B, _GREY)
    for i, s in enumerate(data["standar"]):
        r = 5 + i
        _put(wn, f"B{r}", i + 1)
        _put(wn, f"C{r}", s["kode"], F_B, al=_LEFT)
        _put(wn, f"D{r}", s["judul"], F_N, al=_LEFT)
        _put(wn, f"E{r}", s["peran"], F_N, al=_LEFT)
        _put(wn, f"F{r}", s["catatan"], F_N, al=_LEFT)
        wn.row_dimensions[r].height = 42
    _put(wn, f"B{5 + len(data['standar']) + 1}", data["catatan_standar"], F_I, al=_NW, border=False)
    for col, w in {"A": 2, "B": 5, "C": 30, "D": 44, "E": 52, "F": 44}.items():
        wn.column_dimensions[col].width = w
    wn.page_setup.orientation = "landscape"
    wn.sheet_properties.pageSetUpPr.fitToPage = True
    wn.page_setup.fitToWidth, wn.page_setup.fitToHeight = 1, 0

    # ------------------------------------------------ sheet utama
    for r, (lab, key) in enumerate([("PEKERJAAN :", "pekerjaan"), ("LOKASI :", "lokasi"),
                                    ("TAHUN :", "tahun"), ("ITEM PEKERJAAN :", "item")], start=2):
        merge(f"B{r}:C{r}", lab, font=F_B, al=_NW, border=False)
        merge(f"D{r}:G{r}", f": {proyek.get(key, '')}", font=F_N, al=_NW, border=False)
    merge("B7:G7", "Landasan — Plumbing: SNI 8153:2015 · Tabel pemakaian air: SNI 03-7065-2005 · Tangki septik: SNI 2398:2017 · "
                   "Baku mutu: Permen LHK P.68/2016 · SPALD: Permen PUPR 4/2017 (lengkap pada sheet 'SNI & Peraturan')",
          font=F_N, al=_LEFT, border=False)
    ws.row_dimensions[7].height = 30

    r = 9

    def judul(t):
        nonlocal r
        put(f"B{r}", t, F_B, al=_NW, border=False)
        r += 1

    # ---- 1. kebutuhan air bersih
    judul("1. Kebutuhan Air Bersih")
    for col, t in zip("BCDEFG", ["No", "Fungsi Bangunan", "Jumlah", "Satuan", "Standar (L/satuan/hari)", "Q (L/hari)"]):
        put(f"{col}{r}", t, F_B, _GREY)
    r += 1
    r1 = r
    for i, x in enumerate(h["rinc"].itertuples(), start=1):
        put(f"B{r}", i)
        put(f"C{r}", x.Fungsi, F_IN, al=_LEFT)
        put(f"D{r}", x.Jumlah, F_IN, fmt="#,##0.##")
        put(f"E{r}", f"=VLOOKUP(C{r},{TAB},4,FALSE)", al=_LEFT)
        put(f"F{r}", f"=VLOOKUP(C{r},{TAB},2,FALSE)", fmt="#,##0")
        put(f"G{r}", f"=D{r}*F{r}", fmt="#,##0")
        r += 1
    r2 = r - 1
    merge(f"B{r}:F{r}", "Total Q harian (rata-rata)  [L/hari]", font=F_B, fill=_GREY)
    put(f"G{r}", f"=SUM(G{r1}:G{r2})", F_B, _GREY, fmt="#,##0")
    QH = f"$G${r}"
    r += 1
    merge(f"B{r}:F{r}", "Faktor puncak", font=F_N, al=_LEFT)
    put(f"G{r}", P["faktor_puncak"], F_IN, fmt="0%")
    FP = f"$G${r}"
    r += 1
    merge(f"B{r}:F{r}", "Q rencana = Q harian × (1 + faktor puncak)  [L/hari]", font=F_N, al=_LEFT)
    put(f"G{r}", f"={QH}*(1+{FP})", F_N, fmt="#,##0")
    QR = f"$G${r}"
    r += 1
    merge(f"B{r}:F{r}", "Q rencana  [m³/hari]", font=F_N, al=_LEFT)
    put(f"G{r}", f"={QR}/1000", F_N, fmt="0.000")
    QM = f"$G${r}"
    r += 1
    merge(f"B{r}:F{r}", "Pembulatan volume tandon ke atas, kelipatan  [m³]", font=F_N, al=_LEFT)
    put(f"G{r}", P["pembulatan"], F_IN, fmt="0.00")
    STEP = f"$G${r}"
    r += 1
    merge(f"B{r}:F{r}", "VOLUME TANDON RENCANA  [m³]", font=F_B, fill=_GREY)
    put(f"G{r}", f"=ROUNDUP({QM}/{STEP},0)*{STEP}", F_B, _GREY, fmt="0.00")
    VT = f"$G${r}"
    r += 2

    def baris(label, val, unit, rumus, font=F_N, fmt="#,##0.00", fill=None, val_font=None):
        """label C | nilai D | satuan E | rumus F:G"""
        nonlocal r
        put(f"C{r}", label, font, fill, al=_LEFT)
        put(f"D{r}", val, val_font or font, fill, fmt=fmt)
        put(f"E{r}", unit, F_N, fill)
        merge(f"F{r}:G{r}", rumus, font=F_I if not fill else F_B, fill=fill, al=_LEFT)
        ref = f"$D${r}"
        r += 1
        return ref

    # ---- 2. air limbah
    judul("2. Air Limbah (Pembagian Grey Water dan Black Water)")
    dasar_ref = QR if P["dasar"].startswith("Q dengan") else QH
    D = baris("Debit dasar perhitungan", f"={dasar_ref}", "L/hari", P["dasar"], fmt="#,##0.0")
    FL = baris("Faktor air limbah (% dari debit dasar)", P["faktor_limbah"], "%",
               "Excel: 100%; acuan umum 80% bila sebagian air tidak menjadi limbah", fmt="0%", val_font=F_IN)
    FB = baris("Fraksi black water (kakus)", P["fraksi_black"], "%", "Excel: 20%", fmt="0%", val_font=F_IN)
    LT = baris("Total air limbah", f"={D}*{FL}", "L/hari", "debit dasar × faktor air limbah", F_B, "#,##0.0", _GREY)
    BL = baris("Black water", f"={LT}*{FB}", "L/hari", "total × fraksi black water", fmt="#,##0.0")
    baris("Grey water", f"={LT}-{BL}", "L/hari", "total − black water", fmt="#,##0.0")
    r += 1

    # ---- 3. tangki septik
    judul("3. Air Kotor (Black Water) dan Tangki Septik (SNI 2398:2017)")
    sd = data["septik"]
    N = baris("Jumlah pemakai (n)", h["n"], "orang", "dari tabel penggunaan (penghuni/pemakai)", fmt="0", val_font=F_IN)
    TD = baris("Waktu detensi (td)", P["td"], "hari", None, fmt="0.0", val_font=F_IN)
    ws[f"F{r - 1}"].value = (f'=IF(AND({TD}>={sd["detensi_min"]},{TD}<={sd["detensi_max"]}),'
                             f'"OK (rentang {sd["detensi_min"]}–{sd["detensi_max"]} hari)","Di luar rentang {sd["detensi_min"]}–{sd["detensi_max"]} hari")')
    QL = baris("Produksi lumpur (QL)", P["ql"], "L/orang/tahun", None, fmt="0", val_font=F_IN)
    ws[f"F{r - 1}"].value = (f'=IF(AND({QL}>={sd["lumpur_min"]},{QL}<={sd["lumpur_max"]}),'
                             f'"OK (rentang {sd["lumpur_min"]}–{sd["lumpur_max"]})","Di luar rentang {sd["lumpur_min"]}–{sd["lumpur_max"]}")')
    PP = baris("Periode pengurasan (PP)", P["pp"], "tahun", None, fmt="0", val_font=F_IN)
    ws[f"F{r - 1}"].value = (f'=IF(AND({PP}>={sd["pengurasan_min"]},{PP}<={sd["pengurasan_max"]}),'
                             f'"OK (rentang {sd["pengurasan_min"]}–{sd["pengurasan_max"]} tahun)","Di luar rentang {sd["pengurasan_min"]}–{sd["pengurasan_max"]} tahun")')
    AM = baris("Ambang bebas", P["ambang"], "m", "tinggi ruang bebas di atas muka air", fmt="0.00", val_font=F_IN)
    HA = baris("Kedalaman air efektif", P["h_air"], "m", "pilihan perencana", fmt="0.00", val_font=F_IN)
    RS = baris("Rasio panjang : lebar", P["rasio"], "P : L", "pilihan perencana", fmt="0.0", val_font=F_IN)
    r += 1
    for col, t in zip("CDEFG", ["Uraian", "Tercampur", "Terpisah", "Satuan", "Rumus"]):
        put(f"{col}{r}", t, F_B, _GREY)
    r += 1

    def hasil(label, fd, fe, unit, rumus, fmt="0.000", bold=False):
        nonlocal r
        f = F_B if bold else F_N
        fill = _GREY if bold else None
        put(f"C{r}", label, f, fill, al=_LEFT)
        put(f"D{r}", fd, f, fill, fmt=fmt)
        put(f"E{r}", fe, f, fill, fmt=fmt)
        put(f"F{r}", unit, F_N, fill)
        put(f"G{r}", rumus, F_I if not bold else F_B, fill, al=_LEFT)
        r += 1
        return r - 1

    q = hasil("Debit air limbah masuk", f"={LT}", f"={BL}", "L/hari",
              "tercampur: grey + black; terpisah: black saja", "#,##0.0")
    va = hasil("Volume air", f"=D{q}*{TD}/1000", f"=E{q}*{TD}/1000", "m³", "Q × td / 1000")
    vl = hasil("Volume lumpur", f"={QL}*{N}*{PP}/1000", f"={QL}*{N}*{PP}/1000", "m³", "QL × n × PP / 1000")
    vb = hasil("Volume basah", f"=D{va}+D{vl}", f"=E{va}+E{vl}", "m³", "V air + V lumpur")
    lu = hasil("Luas denah", f"=D{vb}/{HA}", f"=E{vb}/{HA}", "m²", "V basah / kedalaman air")
    le = hasil("Lebar", f"=SQRT(D{lu}/{RS})", f"=SQRT(E{lu}/{RS})", "m", "√(Luas / rasio P:L)", "0.00")
    hasil("Panjang", f"={RS}*D{le}", f"={RS}*E{le}", "m", "rasio × Lebar", "0.00")
    vm = hasil("Volume ambang bebas", f"=D{lu}*{AM}", f"=E{lu}*{AM}", "m³", "Luas × ambang bebas")
    vt = hasil("VOLUME TOTAL TANGKI", f"=D{vb}+D{vm}", f"=E{vb}+E{vm}", "m³", "V basah + V ambang bebas", bold=True)
    hasil("Tinggi total", f"={HA}+{AM}", f"={HA}+{AM}", "m", "kedalaman air + ambang bebas", "0.00")
    put(f"C{r}", f"Cek: black water contoh SNI ({data['air_limbah']['black_water_l_org_hari']} L/orang/hari × n)", F_I, al=_LEFT)
    put(f"D{r}", None)
    put(f"E{r}", f"={data['air_limbah']['black_water_l_org_hari']}*{N}", F_I, fmt="#,##0.0")
    put(f"F{r}", "L/hari", F_I)
    put(f"G{r}", "pembanding terhadap black water pada sistem terpisah", F_I, al=_LEFT)
    r += 2

    # ---- 4. rekap
    judul("4. Rekap")
    for col, t in zip("CDE", ["Keterangan", "Jumlah", "Satuan"]):
        put(f"{col}{r}", t, F_B, _GREY)
    r += 1
    rek = [("Penghuni / pemakai", f"={N}", "orang", "0", False),
           ("Jumlah pemakaian air harian", f"={QH}", "L/hari", "#,##0", False),
           ("Jumlah pemakaian air harian", f"={QH}/1000", "m³/hari", "0.000", False),
           ("Penggunaan peak time air harian", f"={QH}/1000*{FP}", "m³/hari", "0.000", False),
           ("VOLUME TOTAL", f"={QM}", "m³", "0.000", True),
           ("VOLUME TANDON RENCANA", f"={VT}", "m³", "0.00", True),
           ("VOLUME TANGKI SEPTIK (tercampur)", f"=D{vt}", "m³", "0.00", True)]
    for lab, fx, unit, nf, b in rek:
        f = F_B if b else F_N
        fill = _GREY if b else None
        put(f"C{r}", lab, f, fill, al=_LEFT)
        put(f"D{r}", fx, f, fill, fmt=nf)
        put(f"E{r}", unit, F_N, fill)
        r += 1
    r += 1
    put(f"B{r}", "Keterangan: sel biru = input (boleh diubah); sel hitam = rumus otomatis.", F_I, al=_NW, border=False)
    put(f"B{r + 1}", f"Sumber tabel pemakaian air: {data['sumber_tabel']}.", F_I, al=_NW, border=False)
    put(f"B{r + 2}", data["catatan_standar"], F_I, al=_NW, border=False)

    for col, w in {"A": 2, "B": 5, "C": 50, "D": 15, "E": 17, "F": 22, "G": 50}.items():
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# -------------------------------------------------------------------- PDF
def _cari_soffice() -> str | None:
    p = shutil.which("soffice") or shutil.which("libreoffice")
    if p:
        return p
    for cand in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                 r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if os.path.exists(cand):
            return cand
    return None


def excel_ke_pdf(xlsx_bytes: bytes) -> bytes | None:
    """xlsx -> pdf via LibreOffice. None kalau LibreOffice tidak terpasang / gagal."""
    exe = _cari_soffice()
    if not exe:
        return None
    with tempfile.TemporaryDirectory() as d:
        src = os.path.join(d, "laporan.xlsx")
        with open(src, "wb") as f:
            f.write(xlsx_bytes)
        try:
            subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", d, src],
                           check=True, timeout=120, capture_output=True)
        except (subprocess.SubprocessError, OSError):
            return None
        out = os.path.join(d, "laporan.pdf")
        if not os.path.exists(out):
            return None
        with open(out, "rb") as f:
            return f.read()
