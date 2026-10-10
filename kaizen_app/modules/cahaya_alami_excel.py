"""Modul Pencahayaan Alami - ekspor Excel (rumus hidup, desain konsisten) dan PDF."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break

from modules.cahaya_alami_denah import denah_png

NAVY, AQUA, SOFT, LINE = "0F3D63", "35B6C9", "EAF3F9", "D5E1EB"
_THIN = Side(style="thin", color=LINE)
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_F_FILL = PatternFill("solid", fgColor=NAVY)
_SOFT = PatternFill("solid", fgColor=SOFT)
_MID = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_NW = Alignment(horizontal="left", vertical="center", wrap_text=False)
F_N = Font(name="Arial", size=10)
F_B = Font(name="Arial", size=10, bold=True)
F_IN = Font(name="Arial", size=10, color="0000FF")      # biru = input
F_H = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F_I = Font(name="Arial", size=9, italic=True, color="6B7C8C")
F_T = Font(name="Arial", size=15, bold=True, color="FFFFFF")
F_S = Font(name="Arial", size=12, bold=True, color=NAVY)
MAIN = "'Cahaya Alami'"


def _put(ws, ref, val, font=F_N, fill=None, al=_MID, fmt=None, border=True):
    c = ws[ref]
    c.value = val
    c.font, c.alignment = font, al
    if border:
        c.border = _BOX
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt


def _merge(ws, rng, val, **kw):
    ws.merge_cells(rng)
    _put(ws, rng.split(":")[0], val, **kw)
    if kw.get("border", True):
        for row in ws[rng]:
            for c in row:
                c.border = _BOX


def _judul(ws, rng, teks, tinggi=32):
    _merge(ws, rng, teks, font=F_T, fill=_F_FILL, al=_NW, border=False)
    ws.row_dimensions[int(rng.split(":")[0][1:])].height = tinggi


def _status_cf(ws, rng, first):
    for teks, warna, fc in [("MEMENUHI", "DDF3E6", "1E8E5A"), ("KURANG", "FCE8E6", "C0392B"),
                            ("TIDAK ADA BUKAAN", "FCE8E6", "C0392B"), ("AREA TERBUKA", "ECEFF2", "6B7C8C")]:
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{first}="{teks}"'],
                                      fill=PatternFill("solid", bgColor=warna, fgColor=warna), font=Font(bold=True, color=fc)))


def build_excel(h: dict, data: dict, proyek: dict) -> bytes:
    t, b, thr = h["tabel"], h["bukaan"], h["ambang"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Cahaya Alami"
    wr = wb.create_sheet("Rekap")
    wd = wb.create_sheet("Denah")
    wn = wb.create_sheet("SNI & Peraturan")
    put = lambda *a, **k: _put(ws, *a, **k)
    merge = lambda *a, **k: _merge(ws, *a, **k)

    # ====================================== SNI & peraturan
    _judul(wn, "B1:F1", "LANDASAN PERENCANAAN: SNI & PERATURAN TERKAIT PENCAHAYAAN ALAMI")
    for col, tx in zip("BCDEF", ["No", "Standar / Peraturan", "Judul", "Peran pada perhitungan", "Catatan"]):
        _put(wn, f"{col}3", tx, F_H, _F_FILL)
    for i, s in enumerate(data["standar"]):
        r = 4 + i
        _put(wn, f"B{r}", i + 1)
        _put(wn, f"C{r}", s["kode"], F_B, al=_LEFT)
        _put(wn, f"D{r}", s["judul"], F_N, al=_LEFT)
        _put(wn, f"E{r}", s["peran"], F_N, al=_LEFT)
        _put(wn, f"F{r}", s["catatan"], F_N, al=_LEFT)
        wn.row_dimensions[r].height = 48
    _put(wn, f"B{5 + len(data['standar'])}", data["catatan_standar"], F_I, al=_NW, border=False)
    for col, w in {"A": 2, "B": 5, "C": 34, "D": 44, "E": 56, "F": 46}.items():
        wn.column_dimensions[col].width = w
    wn.page_setup.orientation = "landscape"
    wn.sheet_properties.pageSetUpPr.fitToPage = True
    wn.page_setup.fitToWidth, wn.page_setup.fitToHeight = 1, 0

    # ====================================== Denah
    _judul(wd, "B1:N1", "PERBANDINGAN LUAS BUKAAN DAN LUAS LANTAI (ILUSTRASI PROPORSIONAL, BUKAN POSISI JENDELA)")
    img = XLImage(BytesIO(denah_png(h)))
    ratio = img.height / max(1, img.width)
    img.width, img.height = 1150, int(1150 * ratio)
    wd.add_image(img, "B3")
    wd.sheet_view.showGridLines = False
    wd.page_setup.orientation = "landscape"
    wd.sheet_properties.pageSetUpPr.fitToPage = True
    wd.page_setup.fitToWidth, wd.page_setup.fitToHeight = 1, 1

    # ====================================== sheet utama
    for col, w in {"A": 2, "B": 5, "C": 30, "D": 17, "E": 10, "F": 10, "G": 10, "H": 11, "I": 11, "J": 11, "K": 12,
                   "L": 11, "M": 12, "N": 20, "O": 14}.items():
        ws.column_dimensions[col].width = w
    _judul(ws, "B1:N1", "PERHITUNGAN PENCAHAYAAN ALAMI PADA BANGUNAN — RASIO LUAS BUKAAN TERHADAP LUAS LANTAI", 34)
    for r, (lab, key) in enumerate([("PEKERJAAN", "pekerjaan"), ("LOKASI", "lokasi"), ("TAHUN", "tahun"),
                                    ("ITEM PEKERJAAN", "item")], start=3):
        merge(f"B{r}:C{r}", lab, font=Font(name="Arial", size=10, bold=True, color=NAVY), fill=_SOFT, al=_NW, border=False)
        merge(f"D{r}:J{r}", f": {proyek.get(key, '')}", font=F_N, fill=_SOFT, al=_NW, border=False)
    merge("B8:N8", "Landasan — Pencahayaan alami: SNI 03-2396-2001 · Persyaratan kesehatan perumahan: Kepmenkes 829/Menkes/SK/VII/1999 "
                   "(lengkap pada sheet 'SNI & Peraturan')", font=F_I, al=_LEFT, border=False)

    r = 10

    def sec(no, judul):
        nonlocal r
        put(f"B{r}", no, F_H, _F_FILL)
        merge(f"C{r}:N{r}", judul, font=F_S, al=_NW, border=False)
        for ci in range(3, 15):
            ws.cell(row=r, column=ci).border = Border(bottom=Side(style="medium", color=AQUA))
        ws.row_dimensions[r].height = 22
        r += 1

    sec(1, "Dasar Perhitungan")
    merge(f"C{r}:H{r}", "Rasio = Aj / Ar  ≥  ambang minimum", font=Font(name="Cambria", size=13, italic=True, color=NAVY),
          fill=_SOFT, al=_MID)
    ws.row_dimensions[r].height = 26
    r += 1
    for a, c in [("Aj", "luas bukaan cahaya / jendela (m²) = Σ lebar × tinggi × jumlah × faktor efektif"),
                 ("Ar", "luas lantai ruangan (m²) = P × L"),
                 ("Aj minimum", "ambang × Ar"), ("Catatan", "rasio Aj/Ar adalah pemeriksaan awal; penilaian formal SNI 03-2396-2001 memakai faktor langit (FL)")]:
        put(f"C{r}", f"{a}  =  {c}", F_I, al=_NW, border=False)
        r += 1
    r += 1

    sec(2, "Parameter")
    put(f"C{r}", "Ambang rasio minimum (Aj/Ar)", al=_LEFT)
    put(f"D{r}", thr, F_IN, fmt="0.0%")
    AMB = f"$D${r}"
    r += 2

    sec(3, "Daftar Bukaan (Aj)")
    for col, tx in zip("BCDEFGHI", ["NO", "NAMA RUANGAN", "JENIS BUKAAN", "Lebar (m)", "Tinggi (m)", "Jumlah", "Faktor efektif", "Aj (m²)"]):
        put(f"{col}{r}", tx, F_H, _F_FILL)
    ws.row_dimensions[r].height = 28
    r += 1
    b0 = r
    for i, x in enumerate(b.itertuples(), start=1):
        put(f"B{r}", i)
        put(f"C{r}", x.Ruangan, F_IN, al=_LEFT)
        put(f"D{r}", x.Jenis, F_IN, al=_LEFT)
        put(f"E{r}", x.W, F_IN, fmt="0.00")
        put(f"F{r}", x.H, F_IN, fmt="0.00")
        put(f"G{r}", x.n, F_IN, fmt="0")
        put(f"H{r}", x.f, F_IN, fmt="0.00")
        put(f"I{r}", f"=E{r}*F{r}*G{r}*H{r}", fmt="0.000")
        r += 1
    b1 = max(b0, r - 1)
    if b.empty:
        put(f"C{r}", "—", al=_LEFT)
        r += 1
        b1 = b0
    merge(f"B{r}:H{r}", "TOTAL Aj", font=F_B, fill=_SOFT, al=_NW)
    put(f"I{r}", f"=SUM(I{b0}:I{b1})", F_B, _SOFT, fmt="0.000")
    r += 2

    ws.row_breaks.append(Break(id=r - 1))
    sec(4, "Perhitungan Rasio per Ruangan")
    heads = ["NO", "NAMA RUANGAN", "JENIS RUANG", "P (m)", "L (m)", "Ar (m²)", "Jumlah bukaan", "Aj (m²)", "Rasio Aj/Ar",
             "Aj minimum (m²)", "Selisih (m²)", "Tambahan (m²)", "KETERANGAN"]
    for col, tx in zip("BCDEFGHIJKLMN", heads):
        put(f"{col}{r}", tx, F_H, _F_FILL)
    ws.row_dimensions[r].height = 30
    r += 1
    first = r
    RB, RJ, RA = f"$C${b0}:$C${b1}", f"$I${b0}:$I${b1}", f"$G${b0}:$G${b1}"
    for lt, grp in t.groupby("Lantai", sort=False):
        merge(f"B{r}:N{r}", lt.upper(), font=F_B, fill=_SOFT, al=_NW)
        r += 1
        for i, x in enumerate(grp.itertuples(), start=1):
            nm = lambda v: None if (v is None or v != v) else v
            put(f"B{r}", i)
            put(f"C{r}", x.Ruangan, F_IN, al=_LEFT)
            put(f"D{r}", x.Jenis, F_IN, al=_LEFT)
            put(f"E{r}", nm(x.P), F_IN, fmt="0.00")
            put(f"F{r}", nm(x.L), F_IN, fmt="0.00")
            put(f"G{r}", f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(F{r})),E{r}*F{r},"")', fmt="0.00")
            put(f"H{r}", f"=SUMIFS({RA},{RB},C{r})", fmt="0")
            put(f"I{r}", f"=SUMIFS({RJ},{RB},C{r})", fmt="0.000")
            put(f"J{r}", f'=IF(OR(D{r}="Area terbuka",G{r}=""),"",I{r}/G{r})', F_B, fmt="0.0%")
            put(f"K{r}", f'=IF(J{r}="","",{AMB}*G{r})', fmt="0.000")
            put(f"L{r}", f'=IF(K{r}="","",I{r}-K{r})', fmt='+0.000;-0.000;0.000')
            put(f"M{r}", f'=IF(L{r}="","",IF(L{r}<0,-L{r},0))', fmt="0.000")
            put(f"N{r}", f'=IF(D{r}="Area terbuka","AREA TERBUKA",IF(G{r}="","",IF(I{r}<=0,"TIDAK ADA BUKAAN",'
                         f'IF(J{r}>={AMB}-0.000000001,"MEMENUHI","KURANG"))))', F_B)
            ws[f"O{r}"] = lt
            r += 1
    last = r - 1
    merge(f"B{r}:F{r}", "TOTAL (ruang tertutup)", font=F_B, fill=_SOFT, al=_NW)
    put(f"G{r}", f'=SUMIFS(G{first}:G{last},D{first}:D{last},"<>Area terbuka")', F_B, _SOFT, fmt="0.00")
    put(f"H{r}", None, F_B, _SOFT)
    put(f"I{r}", f'=SUMIFS(I{first}:I{last},D{first}:D{last},"<>Area terbuka")', F_B, _SOFT, fmt="0.000")
    put(f"J{r}", f'=IF(G{r}>0,I{r}/G{r},"")', F_B, _SOFT, fmt="0.0%")
    for col in "KL":
        put(f"{col}{r}", None, F_B, _SOFT)
    put(f"M{r}", f"=SUM(M{first}:M{last})", F_B, _SOFT, fmt="0.000")
    put(f"N{r}", None, F_B, _SOFT)
    tot = r
    _status_cf(ws, f"N{first}:N{last}", f"N{first}")
    r += 1
    put(f"B{r}", "Area terbuka (garasi, teras, taman) tidak dihitung. Ruang tertutup tanpa bukaan berstatus TIDAK ADA BUKAAN. "
                 "Kolom Tambahan = luas bukaan yang masih harus ditambah agar mencapai ambang.", F_I, al=_NW, border=False)
    r += 2

    sec(5, "Kesimpulan")
    kes = (f'="Pemeriksaan pencahayaan alami dilakukan dengan rasio Aj/Ar dengan ambang minimum "&TEXT({AMB},"0.0%")&". Dari "&'
           f'COUNTIFS(D{first}:D{last},"<>Area terbuka",N{first}:N{last},"<>")&" ruangan tertutup, "&COUNTIF(N{first}:N{last},"MEMENUHI")&" ruangan memenuhi"&'
           f'IF(COUNTIF(N{first}:N{last},"KURANG")+COUNTIF(N{first}:N{last},"TIDAK ADA BUKAAN")>0," dan "&(COUNTIF(N{first}:N{last},"KURANG")+COUNTIF(N{first}:N{last},"TIDAK ADA BUKAAN"))&'
           f'" ruangan belum memenuhi (total tambahan bukaan ± "&TEXT(M{tot},"0.000")&" m²)","")&". Rasio keseluruhan "&TEXT(J{tot},"0.0%")&"."')
    merge(f"C{r}:N{r}", kes, font=F_N, fill=PatternFill("solid", fgColor="F1FAF5"), al=_LEFT)
    ws.row_dimensions[r].height = 44
    r += 2
    sec(6, "Saran Penanganan untuk Bangunan yang Sudah Terbangun")
    kurang = t[t["status"].isin(["KURANG", "TIDAK ADA BUKAAN"])]
    if kurang.empty:
        merge(f"C{r}:N{r}", "Tidak ada ruang tertutup yang memerlukan saran retrofit berdasarkan ambang ini.",
              font=F_N, fill=_SOFT, al=_LEFT)
        ws.row_dimensions[r].height = 30
        r += 1
    else:
        for x in kurang.itertuples():
            merge(f"C{r}:N{r}", f"{x.Ruangan} — {x.status}: {x.saran_utama}",
                  font=F_B, fill=_SOFT, al=_LEFT)
            ws.row_dimensions[r].height = 48
            r += 1
            merge(f"C{r}:N{r}", f"Faktor pendukung: {x.saran_pendukung}",
                  font=F_N, al=_LEFT)
            ws.row_dimensions[r].height = 58
            r += 1
    merge(f"C{r}:N{r}", "Saran adalah opsi awal, bukan jaminan memenuhi standar. Verifikasi faktor langit sesuai SNI 03-2396-2001, "
                         "kondisi panas/silau, privasi, kebocoran, dan keamanan struktur dengan tenaga ahli. Bukaan ke ruang dalam "
                         "atau pantulan cahaya tidak otomatis dihitung sebagai bukaan luar.",
          font=F_I, al=_LEFT, border=False)
    ws.row_dimensions[r].height = 42
    r += 2
    merge(f"B{r}:N{r}", "Keterangan: sel biru = input (boleh diubah); sel hitam = rumus otomatis. " + data["catatan_standar"],
          font=F_I, al=_LEFT, border=False)
    ws.row_dimensions[r].height = 28

    ws.column_dimensions["O"].hidden = True
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0

    # ====================================== Rekap
    lantai = list(h["lantai"]["Lantai"])
    rng = lambda col: f"{MAIN}!${col}${first}:${col}${last}"
    _judul(wr, "B1:J1", "REKAP PENCAHAYAAN ALAMI PER LANTAI")
    for col, tx in zip("BCDEFGHIJ", ["Lantai", "Ruang tertutup", "Area terbuka", "Luas lantai (m²)", "Luas bukaan (m²)",
                                     "Rasio", "Memenuhi", "Kurang", "Tambahan (m²)"]):
        _put(wr, f"{col}3", tx, F_H, _F_FILL)
    wr.row_dimensions[3].height = 28
    rr = 4
    for lt in lantai:
        _put(wr, f"B{rr}", lt, F_B, al=_LEFT)
        _put(wr, f"C{rr}", f'=COUNTIFS({rng("O")},$B{rr},{rng("N")},"<>AREA TERBUKA",{rng("N")},"<>")', fmt="0")
        _put(wr, f"D{rr}", f'=COUNTIFS({rng("O")},$B{rr},{rng("N")},"AREA TERBUKA")', fmt="0")
        _put(wr, f"E{rr}", f'=SUMIFS({rng("G")},{rng("O")},$B{rr},{rng("D")},"<>Area terbuka")', fmt="0.00")
        _put(wr, f"F{rr}", f'=SUMIFS({rng("I")},{rng("O")},$B{rr},{rng("D")},"<>Area terbuka")', fmt="0.000")
        _put(wr, f"G{rr}", f'=IF(E{rr}>0,F{rr}/E{rr},"")', F_B, fmt="0.0%")
        _put(wr, f"H{rr}", f'=COUNTIFS({rng("O")},$B{rr},{rng("N")},"MEMENUHI")', fmt="0")
        _put(wr, f"I{rr}", f'=COUNTIFS({rng("O")},$B{rr},{rng("N")},"KURANG")+COUNTIFS({rng("O")},$B{rr},{rng("N")},"TIDAK ADA BUKAAN")', fmt="0")
        _put(wr, f"J{rr}", f'=SUMIFS({rng("M")},{rng("O")},$B{rr})', fmt="0.000")
        rr += 1
    _put(wr, f"B{rr}", "TOTAL", F_B, _SOFT, al=_LEFT)
    for col, nf in zip("CDEFHIJ", ["0", "0", "0.00", "0.000", "0", "0", "0.000"]):
        _put(wr, f"{col}{rr}", f"=SUM({col}4:{col}{rr - 1})", F_B, _SOFT, fmt=nf)
    _put(wr, f"G{rr}", f'=IF(E{rr}>0,F{rr}/E{rr},"")', F_B, _SOFT, fmt="0.0%")
    for col, w in {"A": 2, "B": 20, "C": 14, "D": 13, "E": 15, "F": 15, "G": 11, "H": 12, "I": 11, "J": 14}.items():
        wr.column_dimensions[col].width = w
    wr.sheet_view.showGridLines = False
    wr.page_setup.orientation = "landscape"
    wr.sheet_properties.pageSetUpPr.fitToPage = True
    wr.page_setup.fitToWidth, wr.page_setup.fitToHeight = 1, 1

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
